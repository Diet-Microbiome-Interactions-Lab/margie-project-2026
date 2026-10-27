"""The Excel workbook (fingerprint/make-final-excel.py) tints each row by CONFIDENCE_TIER, the default
(adjacency) tier, and sizes columns by their unprefixed names, so "Column-M: best_consensus_..."
gets its own width."""

import csv
import importlib.util
import sys
from pathlib import Path

from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / 'bioinformatics_tools' / 'workflow_tools' / 'fingerprint' / 'make-final-excel.py'
spec = importlib.util.spec_from_file_location('make_final_excel_t', SCRIPT)
excel = importlib.util.module_from_spec(spec)
spec.loader.exec_module(excel)

HEADERS = ['Column-A: organism_name', 'Column-M: best_consensus_product_descriptor',
           'Column-AS: CONFIDENCE_TIER', 'Column-AT: CONFIDENCE_TIER_hybrid', 'Column-AU: NEEDS_REVIEW?']


def test_rows_take_the_default_tier_and_columns_their_widths(tmp_path, monkeypatch):
    tsv = tmp_path / 'final.tsv'
    with open(tsv, 'w', newline='') as fh:
        w = csv.writer(fh, delimiter='\t')
        w.writerow(HEADERS)
        w.writerow(['coli', 'glucokinase', 'low', 'highest', 'no'])
        w.writerow(['coli', 'tRNA-Ala', 'NOT_APPLICABLE_NON_CODING', 'NOT_APPLICABLE_NON_CODING', 'n/a'])
    out = tmp_path / 'final.xlsx'
    monkeypatch.setattr(sys, 'argv', ['x', '--input', str(tsv), '--output', str(out)])
    excel.main()
    ws = load_workbook(out)['Annotation Results']
    low = excel._tint_hex(excel._TIER_BRIGHT['low'], excel._ROW_TINT)
    assert ws.cell(2, 1).fill.fgColor.rgb.endswith(low)
    assert ws.cell(3, 1).fill.fgColor.rgb.endswith(excel._ROW_NONCODING_BG)
    assert ws.column_dimensions['A'].width == 35 and ws.column_dimensions['B'].width == 36
    assert ws.column_dimensions['C'].width == 14
    legend = load_workbook(out)['Legend']
    assert legend.cell(3, 1).value == 'ROW COLOUR = confidence tier (CONFIDENCE_TIER)'
