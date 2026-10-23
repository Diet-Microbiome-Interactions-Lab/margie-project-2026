"""
SSH and job management endpoints.

Thin routing layer — delegates to job_store, job_runner, and ssh utilities.

All endpoints (except /health) require a valid Bearer token. The token is
validated by get_current_user(), which returns the user's cluster credentials.
_build_connection() decrypts the stored private key and builds a per-user
SSHConnection for each request.
"""
import io
import json
import logging
import posixpath
import re
import shlex
import threading
import time
import uuid
from datetime import datetime, timezone

import pandas as pd
import yaml
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from openpyxl.styles import Border, Font, PatternFill, Side
from fastapi.responses import Response, StreamingResponse

from dataclasses import asdict

from bioinformatics_tools.api import build_gate, licensing
from bioinformatics_tools.api.auth import decode_token, decrypt_private_key, get_current_user
from bioinformatics_tools.api.models import GenomeSend, SlurmSend
from bioinformatics_tools.api.services import job_history_client, job_runner, read_cache, tool_assets, user_stores
from bioinformatics_tools.api.services.job_store import job_store
from bioinformatics_tools.utilities import ssh_sftp, ssh_slurm
from bioinformatics_tools.utilities.ssh_connection import (
    make_user_connection, provision_in_background, sync_remote_dane_wf, wait_provisioned,
)
from bioinformatics_tools.workflow_tools.workflow_helpers import GENOME_EXTENSIONS
from bioinformatics_tools.workflow_tools.workflow_registry import (
    MARGIE_SB_PHASED_TOOLS,
    WORKFLOWS,
    REQUIRED_SYSTEM_PARAMS,
    resolve_user_paths,
    workflow_path_params,
)

LOGGER = logging.getLogger(__name__)

def _changes_drop_reads(request: Request):
    """Anything here that changes something (not a GET) drops the user's short-lived reads
    (read_cache), before it runs and again after, so the next page reads afresh."""
    if request.method in ("GET", "HEAD", "OPTIONS"):
        yield
        return
    user_id = None
    auth = request.headers.get("authorization", "")
    if auth.lower().startswith("bearer "):
        try:
            user_id = int(decode_token(auth[7:])["sub"])
        except Exception:
            user_id = None
    if user_id is not None:
        read_cache.drop(user_id)
    yield
    if user_id is not None:
        read_cache.drop(user_id)


router = APIRouter(prefix="/v1/ssh", tags=["ssh"], dependencies=[Depends(_changes_drop_reads)])

# compute.cluster_default.* defaults from the same registry as config.yaml.
_CLUSTER_DEFAULTS: dict[str, object] = {
    p['param'].split('.')[-1]: p.get('default')
    for p in REQUIRED_SYSTEM_PARAMS
    if p['param'].startswith('compute.cluster_default.')
}


def _cluster_default(user_config: dict, key: str) -> str | None:
    """Reads compute.cluster_default.<key>, falling back to the registry default.

    A key missing from an older config.yaml takes the registry default; a value
    the user has written always wins.
    """
    value = (user_config.get('compute', {})
                        .get('cluster_default', {})
                        .get(key))
    if value is None or str(value).strip() == '':
        value = _CLUSTER_DEFAULTS.get(key)
    if value is None:
        return None
    text = str(value).strip()
    return text or None

# Workflows visible on the frontend but not yet implemented.
STUB_WORKFLOWS: set[str] = {"custom_microbiome"}

# (job_id, path) -> (mtime, size, total_lines); cleared on API restart.
_line_count_cache: dict[tuple[str, str], tuple[float, int, int]] = {}


def _validate_relative_path(path: str, *, label: str = "file") -> None:
    """Raises HTTPException(400) if a user-supplied path under a job's work_dir attempts traversal."""
    if path and (path.startswith("/") or ".." in path.split("/")):
        raise HTTPException(status_code=400, detail=f"Invalid {label} path")


def _cfg_get(cfg: dict, key: str, default=None):
    """Returns a nested config value by dot-notation key, or default if any segment is missing."""
    value = cfg
    for part in key.split('.'):
        if isinstance(value, dict) and part in value:
            value = value[part]
        else:
            return default
    return value


def _cfg_set(cfg: dict, key: str, value) -> None:
    """Sets a nested config value by dot-notation key, creating parents."""
    parts = key.split('.')
    target = cfg
    for part in parts[:-1]:
        next_value = target.get(part)
        if not isinstance(next_value, dict):
            next_value = {}
            target[part] = next_value
        target = next_value
    target[parts[-1]] = value


def _first_nonempty(*values):
    for value in values:
        if value is None:
            continue
        if str(value).strip() == "":
            continue
        return value
    return None


def _expand_remote_home(path: str, home_dir: str) -> str:
    if path.startswith("~"):
        return path.replace("~", home_dir, 1)
    return path


def _is_user_scoped_db(path: str, username: str) -> bool:
    """Returns True when the path basename starts with the '<username>-' prefix (pre-marker paths)."""
    return posixpath.basename(path).startswith(f"{username}-")


def _owner_marker_path(path: str, *, is_dir: bool) -> str:
    """Returns the companion marker path that stores ownership metadata."""
    if is_dir:
        return f"{path.rstrip('/')}/.margie-owner.json"
    return f"{path}.margie-owner.json"


def _read_owner_marker(conn, path: str, *, is_dir: bool) -> dict | None:
    """Reads a path's ownership marker JSON, or None if absent or invalid."""
    marker = _owner_marker_path(path, is_dir=is_dir)
    cmd = f"if [ -f {shlex.quote(marker)} ]; then cat {shlex.quote(marker)}; fi"
    exit_code, output = _run_remote_check(conn, cmd)
    if exit_code != 0 or not output:
        return None
    try:
        data = json.loads(output)
        return data if isinstance(data, dict) else None
    except json.JSONDecodeError:
        LOGGER.warning("Ignoring invalid ownership marker at %s", marker)
        return None


def _write_owner_marker(current_user: dict, conn, path: str, *, is_dir: bool,
                        source_path: str | None = None) -> None:
    """Writes the ownership marker for a user-scoped promoted path."""
    marker = _owner_marker_path(path, is_dir=is_dir)
    payload = {
        "scope": "user",
        "owner_username": current_user["username"],
        "owner_cluster_username": current_user["cluster_username"],
        "kind": "directory" if is_dir else "file",
        "path": path,
    }
    if source_path:
        payload["source_path"] = source_path
    ssh_sftp.write_remote_text_file(
        marker,
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        connection=conn,
    )


def _classify_path_scope(current_user: dict, conn, path: str, *, is_dir: bool) -> str:
    """Classifies a path as 'user' (owned by this user) or 'shared', marker first.

    Raises HTTPException if the marker belongs to a different user.
    """
    marker = _read_owner_marker(conn, path, is_dir=is_dir)
    if marker is not None:
        owner = marker.get("owner_username")
        scope = marker.get("scope")
        if scope == "user" and owner == current_user["username"]:
            return "user"
        if scope == "user" and owner and owner != current_user["username"]:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Configured path '{path}' is marked as private to user '{owner}'. "
                    "Please select a shared template path or your own private path."
                ),
            )

    # Pre-marker fallback; scratch stores are named after the cluster username.
    if _is_user_scoped_db(path, current_user["username"]):
        return "user"
    cluster_user = current_user.get("cluster_username") or ""
    return "user" if cluster_user and _is_user_scoped_db(path, cluster_user) else "shared"


def _versioned_user_db_path(template_db: str, username: str, version: int) -> str:
    """Builds '<dir>/<username>-<stem>-vN<ext>' from a shared template DB path."""
    directory = posixpath.dirname(template_db)
    filename = posixpath.basename(template_db)
    stem, ext = posixpath.splitext(filename)
    target_name = f"{username}-{stem}-v{version}{ext}"
    return posixpath.join(directory, target_name) if directory else target_name


def _versioned_user_dir_path(template_dir: str, username: str, version: int) -> str:
    """Builds '<dir>/<username>-<name>-vN' from a shared directory path."""
    parent = posixpath.dirname(template_dir.rstrip('/'))
    name = posixpath.basename(template_dir.rstrip('/'))
    target_name = f"{username}-{name}-v{version}"
    return posixpath.join(parent, target_name) if parent else target_name


def _find_existing_user_db_versions(conn, template_db: str, username: str) -> list[int]:
    """Lists version numbers of username-prefixed copies of template_db."""
    directory = posixpath.dirname(template_db)
    if not directory:
        directory = "."
    filename = posixpath.basename(template_db)
    stem, ext = posixpath.splitext(filename)
    prefix = f"{username}-{stem}-v"

    try:
        entries = ssh_sftp.list_remote_dir(directory, connection=conn)
    except FileNotFoundError:
        return []
    except Exception:
        return []

    versions: list[int] = []
    for entry in entries:
        if entry.get("type") != "file":
            continue
        name = entry.get("name") or ""
        if not name.startswith(prefix) or not name.endswith(ext):
            continue
        middle = name[len(prefix):]
        if ext:
            middle = middle[:-len(ext)]
        if middle.isdigit():
            versions.append(int(middle))
    return sorted(versions)


def _find_existing_user_dir_versions(conn, template_dir: str, username: str) -> list[int]:
    """Lists version numbers of username-prefixed copies of a directory."""
    parent = posixpath.dirname(template_dir.rstrip('/'))
    if not parent:
        parent = "."
    base = posixpath.basename(template_dir.rstrip('/'))
    prefix = f"{username}-{base}-v"

    try:
        entries = ssh_sftp.list_remote_dir(parent, connection=conn)
    except FileNotFoundError:
        return []
    except Exception:
        return []

    versions: list[int] = []
    for entry in entries:
        if entry.get("type") != "directory":
            continue
        name = entry.get("name") or ""
        if not name.startswith(prefix):
            continue
        suffix = name[len(prefix):]
        if suffix.isdigit():
            versions.append(int(suffix))
    return sorted(versions)


def _promote_shared_file_to_user_file(current_user: dict, conn, raw_path: str) -> tuple[str, bool]:
    """Resolves a writable per-user file path, copying the shared template on first use."""
    expanded = _expand_remote_home(raw_path, current_user["home_dir"])
    username = current_user["username"]

    if _classify_path_scope(current_user, conn, expanded, is_dir=False) == "user":
        # Backfills marker metadata for user-prefixed paths that predate markers.
        if _read_owner_marker(conn, expanded, is_dir=False) is None:
            _write_owner_marker(current_user, conn, expanded, is_dir=False)
        return expanded, False

    try:
        ssh_sftp.check_remote_file(expanded, connection=conn)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Shared file path does not exist or is unreadable: '{expanded}'. Details: {exc}",
        )

    existing_versions = _find_existing_user_db_versions(conn, expanded, username)
    if existing_versions:
        target = _versioned_user_db_path(expanded, username, existing_versions[-1])
        if _read_owner_marker(conn, target, is_dir=False) is None:
            _write_owner_marker(current_user, conn, target, is_dir=False, source_path=expanded)
    else:
        target = _versioned_user_db_path(expanded, username, 1)
        target_dir = posixpath.dirname(target) or "."
        cmd = (
            f"mkdir -p {shlex.quote(target_dir)} && "
            f"cp {shlex.quote(expanded)} {shlex.quote(target)}"
        )
        exit_code, output = _run_remote_check(conn, cmd)
        if exit_code != 0:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Could not create user-specific file from shared template. "
                    f"Source: '{expanded}', target: '{target}'. Details: {output or 'copy failed'}"
                ),
            )
        LOGGER.info("Created user-specific shared file copy for %s: %s", username, target)
        _write_owner_marker(current_user, conn, target, is_dir=False, source_path=expanded)
    return target, True


def _promote_shared_dir_to_user_dir(current_user: dict, conn, raw_path: str) -> tuple[str, bool]:
    """Resolves a writable per-user directory path, creating a versioned copy on first use."""
    expanded = _expand_remote_home(raw_path, current_user["home_dir"])
    username = current_user["username"]

    if _classify_path_scope(current_user, conn, expanded.rstrip('/'), is_dir=True) == "user":
        cmd = f"mkdir -p {shlex.quote(expanded)}"
        exit_code, output = _run_remote_check(conn, cmd)
        if exit_code != 0:
            raise HTTPException(
                status_code=400,
                detail=f"Could not ensure user directory exists: '{expanded}'. Details: {output or 'mkdir failed'}",
            )
        if _read_owner_marker(conn, expanded, is_dir=True) is None:
            _write_owner_marker(current_user, conn, expanded, is_dir=True)
        return expanded, False

    existing_versions = _find_existing_user_dir_versions(conn, expanded, username)
    if existing_versions:
        target = _versioned_user_dir_path(expanded, username, existing_versions[-1])
    else:
        target = _versioned_user_dir_path(expanded, username, 1)

    cmd = f"mkdir -p {shlex.quote(target)}"
    exit_code, output = _run_remote_check(conn, cmd)
    if exit_code != 0:
        raise HTTPException(
            status_code=400,
            detail=f"Could not create user-specific directory: '{target}'. Details: {output or 'mkdir failed'}",
        )
    if not existing_versions:
        LOGGER.info("Created user-specific shared directory for %s: %s", username, target)
    if _read_owner_marker(conn, target, is_dir=True) is None:
        _write_owner_marker(current_user, conn, target, is_dir=True, source_path=expanded)
    return target, True


def _promote_shared_main_db_to_user_db(current_user: dict, user_config: dict, conn) -> tuple[str, bool]:
    """Resolves a writable per-user main_database path.

    A username-prefixed path is kept; otherwise the path is treated as a shared
    template and replaced by the highest existing user copy, or a new v1.

    Returns (resolved_main_db_path, config_changed).
    """
    raw = user_config.get("main_database")
    if not raw or str(raw).strip() == "":
        raise HTTPException(
            status_code=400,
            detail="main_database is not configured. Please configure it in your Profile settings.",
        )

    username = current_user["username"]
    expanded = _expand_remote_home(str(raw).strip(), current_user["home_dir"])
    if _classify_path_scope(current_user, conn, expanded, is_dir=False) == "user":
        if _read_owner_marker(conn, expanded, is_dir=False) is None:
            _write_owner_marker(current_user, conn, expanded, is_dir=False)
        return expanded, False

    # Shared template: switches to a user-specific path in the same directory.
    try:
        ssh_sftp.check_remote_file(expanded, connection=conn)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Configured main_database does not exist or is unreadable: '{expanded}'. Details: {exc}",
        )

    existing_versions = _find_existing_user_db_versions(conn, expanded, username)
    if existing_versions:
        target = _versioned_user_db_path(expanded, username, existing_versions[-1])
        if _read_owner_marker(conn, target, is_dir=False) is None:
            _write_owner_marker(current_user, conn, target, is_dir=False, source_path=expanded)
    else:
        target = _versioned_user_db_path(expanded, username, 1)
        target_dir = posixpath.dirname(target) or "."
        cmd = (
            f"mkdir -p {shlex.quote(target_dir)} && "
            f"cp {shlex.quote(expanded)} {shlex.quote(target)}"
        )
        exit_code, output = _run_remote_check(conn, cmd)
        if exit_code != 0:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Could not create user-specific main_database from shared template. "
                    f"Source: '{expanded}', target: '{target}'. Details: {output or 'copy failed'}"
                ),
            )
        LOGGER.info("Created user-specific main_database for %s: %s", username, target)
        _write_owner_marker(current_user, conn, target, is_dir=False, source_path=expanded)

    user_config["main_database"] = target
    return target, True


def _resolve_effective_main_db(current_user: dict, conn, user_config: dict, *, persist: bool) -> str:
    """Resolves main_database, optionally persisting the promoted per-user path to the config."""
    main_db, changed = _promote_shared_main_db_to_user_db(current_user, user_config, conn)
    if changed and persist:
        ssh_sftp.write_remote_yaml(_config_path(current_user["home_dir"]), user_config, connection=conn)
    return main_db


def _run_remote_check(conn, command: str) -> tuple[int, str]:
    ssh = conn.connect()
    _, stdout, stderr = ssh.exec_command(command)
    exit_code = stdout.channel.recv_exit_status()
    output = (stdout.read().decode(errors="replace") + stderr.read().decode(errors="replace")).strip()
    return exit_code, output


def _assert_remote_writable(conn, path: str, *, label: str, treat_as_file: bool = False) -> None:
    target_dir = posixpath.dirname(path) if treat_as_file else path
    if not target_dir:
        target_dir = "."
    probe = posixpath.join(target_dir, f".margie_write_test_{uuid.uuid4().hex}")
    command = (
        f"mkdir -p {shlex.quote(target_dir)} && "
        f"test -w {shlex.quote(target_dir)} && "
        f"touch {shlex.quote(probe)} && rm -f {shlex.quote(probe)}"
    )
    exit_code, output = _run_remote_check(conn, command)
    if exit_code != 0:
        raise HTTPException(
            status_code=400,
            detail=f"{label} is not writable or cannot be created at '{path}'. "
                   f"Please update your Profile config path settings. Details: {output or 'permission/path check failed'}",
        )


def _validate_margie_sb_shared_paths(user_config: dict, conn, home_dir: str) -> None:
    """Fails early if the MARGIE_SB shared storage paths are not writable.

    Namespaced keys are preferred; top-level keys are still read as a fallback.
    """
    writable_paths = [
        (
            _first_nonempty(
                _cfg_get(user_config, 'margie_sb.operon_database.occ_reference_pkl'),
                _cfg_get(user_config, 'operon_database.occ_reference_pkl'),
            ),
            'margie_sb.operon_database.occ_reference_pkl',
            True,
        ),
        (
            _first_nonempty(
                _cfg_get(user_config, 'margie_sb.fingerprint_database.path'),
                _cfg_get(user_config, 'fingerprint_database.path'),
            ),
            'margie_sb.fingerprint_database.path',
            True,
        ),
        (
            _first_nonempty(
                _cfg_get(user_config, 'margie_sb.genome_pool.path'),
                _cfg_get(user_config, 'genome_pool.path'),
            ),
            'margie_sb.genome_pool.path',
            False,
        ),
        (
            _first_nonempty(
                _cfg_get(user_config, 'margie_sb.scoring_results_historical.path'),
                _cfg_get(user_config, 'scoring_results_historical.path'),
            ),
            'margie_sb.scoring_results_historical.path',
            False,
        ),
        (
            _first_nonempty(
                _cfg_get(user_config, 'margie_sb.final_tables_depot.path'),
                _cfg_get(user_config, 'final_tables_depot.path'),
            ),
            'margie_sb.final_tables_depot.path',
            False,
        ),
        (
            _first_nonempty(
                _cfg_get(user_config, 'margie_sb.sqlite_pipeline_snapshot.path'),
                _cfg_get(user_config, 'sqlite_pipeline_snapshot.path'),
            ),
            'margie_sb.sqlite_pipeline_snapshot.path',
            False,
        ),
    ]

    for raw_path, key_name, treat_as_file in writable_paths:
        # Unset means nothing to check: user_stores sets every store before a run.
        if not raw_path:
            continue
        expanded = _expand_remote_home(str(raw_path), home_dir)
        _assert_remote_writable(conn, expanded, label=f"Shared path '{key_name}'", treat_as_file=treat_as_file)


def _resolve_job_work_dir(job_id: str, current_user: dict, conn) -> str:
    """Resolves a job's work_dir from job_store, falling back to persistent history.

    Raises HTTPException(404) if the job is unknown, or 400 if it has no work_dir yet.
    """
    job = job_store.get(job_id)
    if job is not None:
        if job.get("user_id") != current_user["user_id"]:
            raise HTTPException(status_code=403, detail="Access denied")
        work_dir = job.get("work_dir")
    else:
        try:
            user_config = ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn)
        except Exception:
            raise HTTPException(status_code=404, detail="Job not found")

        main_db = user_config.get('main_database')
        row = (
            job_history_client.get_job(
                conn,
                main_db,
                job_id,
                owner_username=current_user["username"],
                owner_cluster_username=current_user["cluster_username"],
            )
            if main_db else None
        )
        if row is None:
            raise HTTPException(status_code=404, detail="Job not found")
        work_dir = row.get("work_dir")

    if not work_dir:
        raise HTTPException(status_code=400, detail="No working directory available for this job")
    return work_dir


_SCORE_TIER_COLORS: dict[str, str] = {
    "highest":          "1a9641",
    "high":             "a6d96a",
    "moderate":         "ffffbf",
    "fair":             "fdae61",
    "low":              "d7191c",
}

_CONFIDENCE_TIER_COLORS: dict[str, str] = {
    "high":             "1a9641",
    "moderate":         "ffffbf",
    "low":              "d7191c",
    "flagged_for_review": "9e1985",
}

# ACS tier: yellow (low) → dark blue (highest), matching make-final-excel.py
_ACS_TIER_ROW_COLORS: dict[str, str] = {
    "low":      "FFFDE7",
    "fair":     "FFF3E0",
    "moderate": "E8F4FC",
    "high":     "D6E8F7",
    "highest":  "C5DDEF",
    "NOT_APPLICABLE_NON_CODING": "F5F5F5",
}

# White text on dark backgrounds, black text on light ones.
_TIER_FONT_COLORS: dict[str, str] = {
    "1a9641": "FFFFFF",
    "a6d96a": "000000",
    "ffffbf": "000000",
    "fdae61": "000000",
    "d7191c": "FFFFFF",
    "9e1985": "FFFFFF",
    "FFFDE7": "000000",
    "FFF3E0": "000000",
    "E8F4FC": "000000",
    "D6E8F7": "000000",
    "C5DDEF": "000000",
    "F5F5F5": "888888",
}


# FINAL rows are tinted by CONFIDENCE_TIER_HYBRID; review rows get a border.
_TIER_BRIGHT = {
    "highest": "1F77FF",   # blue
    "high":    "00B84D",   # green
    "medium":  "FFCC00",   # yellow
    "fair":    "FF8C00",   # orange
    "low":     "EE2233",   # red
}
_ROW_NONCODING_TINT = "F2F2F2"
_ROW_NONCODING_FG = "8A8A8A"


def _tint_hex(h: str, toward_white: float = 0.86) -> str:
    h = h.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    r = round(r + (255 - r) * toward_white)
    g = round(g + (255 - g) * toward_white)
    b = round(b + (255 - b) * toward_white)
    return f"{r:02X}{g:02X}{b:02X}"


def _norm_col(name: str) -> str:
    # Strips header prefixes such as "[AN]-" or "Column-AN: ".
    return re.sub(r"^(?:\[[A-Z]+\]-|Column-[A-Z]+:\s*)", "", str(name or "").strip(), flags=re.IGNORECASE).strip().lower()


def _series_get(row: pd.Series, *candidate_names: str) -> str:
    targets = {_norm_col(n) for n in candidate_names}
    for key in row.index:
        if _norm_col(key) in targets:
            return row.get(key, "")
    return ""


def _has_any_column(df: pd.DataFrame, *candidate_names: str) -> bool:
    targets = {_norm_col(n) for n in candidate_names}
    for col in df.columns:
        if _norm_col(col) in targets:
            return True
    return False


_ROW_TINT = 0.72                                     # tier-colour lightness per row
_REVIEW_SIDE = Side(style="medium", color="000000")  # box border on review rows


def _row_tint(row: pd.Series) -> tuple[str, str]:
    """Returns the (bg, fg) row colour: the tinted CONFIDENCE_TIER_HYBRID colour,
    or grey for rows with no scored tier (empty or NOT_APPLICABLE_NON_CODING)."""
    tier = str(_series_get(row, "confidence_tier_hybrid", "CONFIDENCE_TIER_hybrid")).strip().lower()
    if tier not in _TIER_BRIGHT:
        return _ROW_NONCODING_TINT, _ROW_NONCODING_FG
    return _tint_hex(_TIER_BRIGHT[tier], _ROW_TINT), "000000"


def _apply_review_flag_colors(ws, df: pd.DataFrame) -> None:
    """Tints each data row with its tier colour and boxes rows flagged NEEDS_REVIEW? = yes."""
    n_cols = len(df.columns)
    for offset, (_, row) in enumerate(df.iterrows()):
        r = offset + 2  # row 1 is the header
        bg, fg = _row_tint(row)
        fill = PatternFill(fill_type="solid", fgColor=bg)
        font = Font(color=fg)
        review = str(_series_get(row, "needs_review?", "NEEDS_REVIEW?", "needs_review")).strip().lower() == "yes"
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = fill
            cell.font = font
            if review:                          # box border around the whole review row
                cell.border = Border(
                    top=_REVIEW_SIDE, bottom=_REVIEW_SIDE,
                    left=_REVIEW_SIDE if col == 1 else None,
                    right=_REVIEW_SIDE if col == n_cols else None)


def _apply_tier_row_colors(ws, df: pd.DataFrame) -> None:
    """Colours data rows: review-flag colouring for the FINAL file, else per-tier
    colouring when a tier column exists."""
    if _has_any_column(df, "final_confidence_operon_context", "ADJUSTED_CONFIDENCE_WITH_OPERON_CONTEXT"):
        _apply_review_flag_colors(ws, df)
        return
    if "ACS_tier" in df.columns:
        tier_col, color_map = "ACS_tier", _ACS_TIER_ROW_COLORS
    elif "confidence_score_tier" in df.columns:
        tier_col, color_map = "confidence_score_tier", _SCORE_TIER_COLORS
    elif "confidence_tier" in df.columns:
        tier_col, color_map = "confidence_tier", _CONFIDENCE_TIER_COLORS
    else:
        return
    n_cols = len(df.columns)
    for row_idx, tier_val in enumerate(df[tier_col], start=2):  # row 1 is the header
        hex_color = color_map.get(str(tier_val).strip())
        if hex_color is None:
            continue
        fill = PatternFill(fill_type="solid", fgColor=hex_color)
        font_color = _TIER_FONT_COLORS.get(hex_color, "000000")
        font = Font(color=font_color)
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.fill = fill
            cell.font = font


def _detect_delimiter(path: str, header: str) -> str:
    """Picks a column delimiter from the extension, else by sniffing the header.

    Quoted fields containing delimiters are not handled.
    """
    lower = path.lower()
    if lower.endswith(".csv"):
        return ","
    if lower.endswith(".tsv"):
        return "\t"
    return "\t" if "\t" in header else ","


def _get_available_workflows(cluster_username: str | None = None) -> list[dict]:
    """
    Build the list of available workflows from WORKFLOWS registry.
    Returns detailed metadata for each workflow including tools, params, etc.
    Automatically merges REQUIRED_SYSTEM_PARAMS with workflow-specific params.
    """
    workflows = []

    # Add workflows from WORKFLOWS registry
    for wf_id, wf_key in WORKFLOWS.items():
        # Skip internal test workflows
        if wf_id in ['example', 'selftest']:
            continue

        # Convert dataclass to dict and add computed fields
        wf_dict = asdict(wf_key)
        wf_dict['id'] = wf_key.cmd_identifier
        wf_dict['containers'] = [{'name': sif[0], 'version': sif[1]} for sif in wf_key.sif_files]

        # Merges system, root-path and workflow params, in that order.
        path_params = workflow_path_params(
            wf_id,
            include_sif=wf_key.local_sif_only,
            include_db_root=wf_key.supports_db_root,
            supports_batch_input=wf_key.supports_batch_input,
        )
        wf_dict['configurable_params'] = resolve_user_paths(
            REQUIRED_SYSTEM_PARAMS + path_params + (wf_key.configurable_params or []),
            cluster_username)

        workflows.append(wf_dict)

    # Even stub workflows get system params since they'll need them when implemented
    workflows.append({
        'id': 'custom_microbiome',
        'label': 'Custom Microbiome',
        'description': 'Custom microbiome annotation workflow (coming soon)',
        'full_description': 'A specialized workflow for microbiome annotation. This workflow is currently under development.',
        'tools': [],
        'configurable_params': REQUIRED_SYSTEM_PARAMS,  # Stub still needs system params
        'database_deps': [],
        'docs_url': None,
        'containers': [],
        'cmd_identifier': 'custom_microbiome',
        'snakemake_file': '',
        'other': [],
        'sif_files': [],
    })

    return workflows


def _build_connection(current_user: dict):
    """Decrypt the user's stored private key and return a ready SSHConnection."""
    private_key = decrypt_private_key(current_user['private_key_encrypted'])
    conn = make_user_connection(
        current_user['cluster_host'],
        current_user['cluster_username'],
        private_key,
    )
    # The account's own clone of MARGIE: made, or moved over from an older checkout, while the
    # user looks around, so a run never waits on it unless it starts at once.
    provision_in_background(conn)
    return conn


def _config_path(home_dir: str) -> str:
    """Remote path to the user's BSP config file."""
    return f'{home_dir}/.config/bioinformatics-tools/config.yaml'


def _yaml_block(data: dict, indent: int = 0) -> str:
    block = yaml.safe_dump(data, sort_keys=False, default_flow_style=False).rstrip()
    prefix = ' ' * indent
    return '\n'.join(f'{prefix}{line}' if line else '' for line in block.splitlines())


def _set_nested_value(target: dict, parts: list[str], value):
    current = target
    for part in parts[:-1]:
        current = current.setdefault(part, {})
    current[parts[-1]] = value


def _ordered_workflow_params(workflow_id: str, params: list[dict]) -> list[dict]:
    if workflow_id != 'margie_sb':
        return list(params)

    shared_group_order = {
        'operon_database': 0,
        'fingerprint_database': 1,
        'genome_pool': 2,
        'scoring_results_historical': 3,
        'final_tables_depot': 4,
        'sqlite_pipeline_snapshot': 5,
        'report_figures': 6,
    }
    tool_phase_order = {tool['key']: (tool['phase'], index) for index, tool in enumerate(MARGIE_SB_PHASED_TOOLS)}

    def phase_leaf_order(leaf: str) -> int:
        return {
            'partition': 0,
            'max_parallel_genomes': 1,
            'max_parallel_tools': 2,
            'threads': 3,
            'mem_mb': 4,
            'runtime': 5,
            'db': 6,
            'sif': 7,
        }.get(leaf, 99)

    def sort_key(param: dict) -> tuple:
        parts = param['param'].split('.')
        if len(parts) == 2 and parts[1] in {'default_threads', 'default_mem_mb', 'default_runtime'}:
            return (0, 0, 0, parts[1])
        if len(parts) >= 3 and parts[1].startswith('phase'):
            phase_num = int(parts[1][5:]) if parts[1][5:].isdigit() else 999
            return (1, phase_num, phase_leaf_order(parts[-1]), param['param'])
        if len(parts) >= 3 and parts[1] in shared_group_order:
            return (2, shared_group_order[parts[1]], phase_leaf_order(parts[-1]), param['param'])
        if len(parts) >= 2 and parts[0] == 'margie_sb':
            phase_num, tool_index = tool_phase_order.get(parts[1], (999, 999))
            return (3, phase_num, tool_index, phase_leaf_order(parts[-1]), param['param'])
        return (4, param['param'])

    return sorted(params, key=sort_key)


def _default_params_for_workflow(workflow_id: str, workflow) -> list[dict]:
    """Returns the per-workflow params to materialise in a default config.

    Matches the /workflows metadata shown in Profile: root-path params by
    workflow capability, merged with the workflow's own params.
    """
    path_params = workflow_path_params(
        workflow_id,
        include_sif=workflow.local_sif_only,
        include_db_root=workflow.supports_db_root,
        supports_batch_input=workflow.supports_batch_input,
    )
    combined = path_params + (workflow.configurable_params or [])

    seen: set[str] = set()
    unique: list[dict] = []
    for param in combined:
        key = param.get('param')
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(param)
    return unique


def _build_default_config_payload(cluster_username: str | None = None) -> dict:
    """Builds the default config.yaml contents.

    cluster_username scopes the writable stores to this user (see
    workflow_registry.resolve_user_paths); None keeps the shared paths.
    """
    # Build default config from REQUIRED_SYSTEM_PARAMS
    config: dict = {
        'main_database': '~/.local/share/bioinformatics-tools/my-db.db',
        'compute': {'cluster_default': {}},
    }

    # Populate compute.cluster_default with all defaults from REQUIRED_SYSTEM_PARAMS
    for param in resolve_user_paths(REQUIRED_SYSTEM_PARAMS, cluster_username):
        if param['param'].startswith('compute.cluster_default.'):
            key = param['param'].split('.')[-1]
            default_value = param.get('default')
            # Use empty string for required fields with no default, otherwise use the default
            config['compute']['cluster_default'][key] = default_value if default_value is not None else ''

    for workflow_id, workflow in WORKFLOWS.items():
        params = _default_params_for_workflow(workflow_id, workflow)
        if not params:
            continue

        section: dict = {}
        for param in resolve_user_paths(_ordered_workflow_params(workflow_id, params), cluster_username):
            parts = param['param'].split('.')

            # Namespaced params ("margie_sb.sif_path") nest under the workflow block.
            if parts and parts[0] == workflow_id:
                parts = parts[1:]
            if not parts:
                continue

            default_value = param.get('default')
            if default_value is not None:
                _set_nested_value(section, parts, default_value)

        if section:
            config[workflow_id] = section

    return config


def _build_default_config_text(config: dict) -> str:
    return yaml.safe_dump(config, sort_keys=False, default_flow_style=False)


@router.get("/workflows")
def list_workflows(current_user: dict = Depends(get_current_user)):
    """Returns the user-facing workflows with metadata, path defaults resolved for this user."""
    return _get_available_workflows(current_user.get("cluster_username"))


@router.get("/health")
def health_check():
    """Test endpoint to verify API is working. No auth required."""
    return {"status": "success"}


# What the cluster offers the user, from standard SLURM commands, in one round trip: where the
# site's commands run, the user's accounts and their limits (the account's own association,
# e.g. GrpTRES), what the account's running jobs hold, the QOS limits, the user's jobs, and
# every partition with its nodes.
CLUSTER_SCRIPT = r"""
echo '@@host'; hostname -f 2>/dev/null || hostname
echo '@@assoc'; sacctmgr -nP show assoc user="$USER" format=Account,Partition,QOS,MaxJobs,MaxSubmit 2>/dev/null
accts=$(sacctmgr -nP show assoc user="$USER" format=Account 2>/dev/null | sort -u)
echo '@@accounts'; for a in $accts; do sacctmgr -nP show assoc account="$a" format=Account,User,GrpTRES,GrpJobs 2>/dev/null | awk -F'|' '$2==""'; done
echo '@@qos'; for q in $(sacctmgr -nP show assoc user="$USER" format=QOS 2>/dev/null | tr ',' '\n' | sort -u); do sacctmgr -nP show qos "$q" format=Name,GrpTRES,MaxTRESPU,MaxJobsPU,MaxWall 2>/dev/null; done
echo '@@usage'; for a in $accts; do squeue -h -A "$a" -t R -O 'UserName:40,tres-alloc:400' 2>/dev/null | awk -v a="$a" '{print a "|" $1 "|" $2}'; done
echo '@@jobs'; squeue -h -u "$USER" -o '%i|%j|%P|%T|%M|%l|%D|%C|%N|%r|%a|%o' 2>/dev/null
echo '@@nodes'; sinfo -h -N -o '%P|%n|%T|%C|%m|%e|%G' 2>/dev/null
echo '@@partitions'; sinfo -h -o '%P|%a|%l|%D|%C|%m|%G' 2>/dev/null
echo '@@end'
"""

# Job names and commands that mean a shell on a compute node (sinteractive, srun --pty bash, ...);
# not "(null)", which batch jobs (sbatch --wrap) show as their command too.
_INTERACTIVE = {"interactive", "sinteractive", "bash", "sh", "zsh", "/bin/bash", "/bin/sh", "/bin/zsh", "jupyter", "rstudio"}


def _tres(text: str) -> dict[str, float]:
    """Parses a TRES list ('cpu=8,mem=32G,node=1,gres/gpu=1') into numbers; memory in MB."""
    out: dict[str, float] = {}
    for part in (text or "").split(","):
        key, _, value = part.strip().partition("=")
        if not key or not value:
            continue
        unit = value[-1].upper() if value[-1].isalpha() else ""
        try:
            number = float(value[:-1] if unit else value)
        except ValueError:
            continue
        if key == "mem" or unit in "KMGT" and unit:
            number *= {"K": 1 / 1024, "M": 1, "G": 1024, "T": 1024 * 1024, "": 1}.get(unit, 1)
        out[key] = out.get(key, 0) + number
    return out


def _cpus(field: str) -> dict[str, int]:
    """Parses sinfo's CPUs as allocated/idle/other/total."""
    parts = (field.split("/") + ["0"] * 4)[:4]
    nums = [int(p) if p.isdigit() else 0 for p in parts]
    return dict(zip(("allocated", "idle", "other", "total"), nums))


def parse_cluster(out: str, user: str) -> dict:
    """Turns CLUSTER_SCRIPT's output into the Cluster page's figures."""
    sections: dict[str, list[str]] = {}
    name = ""
    for line in out.splitlines():
        if line.startswith("@@"):
            name = line[2:].strip()
            sections.setdefault(name, [])
        elif name and line.strip():
            sections[name].append(line.rstrip())
    rows = lambda key: [line.split("|") for line in sections.get(key, [])]  # noqa: E731

    host = (sections.get("host") or [""])[0].strip()
    jobs = []
    for f in rows("jobs"):
        if len(f) < 12:
            continue
        job = dict(zip(("id", "name", "partition", "state", "time", "limit", "nodes", "cpus", "nodelist", "reason", "account", "command"), f))
        job["interactive"] = job["name"].strip().lower() in _INTERACTIVE or job["command"].strip().lower() in _INTERACTIVE
        jobs.append(job)

    usage: dict[str, dict] = {}
    for f in rows("usage"):
        if len(f) < 3:
            continue
        acct, who, tres = f[0], f[1].strip(), _tres(f[2])
        u = usage.setdefault(acct, {"all": {}, "mine": {}, "jobs": 0, "my_jobs": 0})
        u["jobs"] += 1
        for k, v in tres.items():
            u["all"][k] = u["all"].get(k, 0) + v
        if who == user:
            u["my_jobs"] += 1
            for k, v in tres.items():
                u["mine"][k] = u["mine"].get(k, 0) + v
    limits = {f[0]: {"tres": _tres(f[2]) if len(f) > 2 else {}, "jobs": f[3] if len(f) > 3 else ""} for f in rows("accounts") if f and f[0]}
    assoc: dict[str, dict] = {}
    for f in rows("assoc"):
        if not f or not f[0]:
            continue
        a = assoc.setdefault(f[0], {"partitions": set(), "qos": set(), "max_jobs": "", "max_submit": ""})
        if len(f) > 1 and f[1]:
            a["partitions"].add(f[1])
        if len(f) > 2:
            a["qos"].update(q for q in f[2].split(",") if q)
        if len(f) > 4:
            a["max_jobs"], a["max_submit"] = f[3], f[4]
    accounts = []
    for acct in sorted(assoc):
        a, lim, use = assoc[acct], limits.get(acct, {}), usage.get(acct, {"all": {}, "mine": {}, "jobs": 0, "my_jobs": 0})
        accounts.append({
            "name": acct,
            "partitions": sorted(a["partitions"]),  # empty: every partition
            "qos": sorted(a["qos"]),
            "max_jobs": a["max_jobs"], "max_submit": a["max_submit"],
            "limits": lim.get("tres", {}), "group_jobs": lim.get("jobs", ""),
            "in_use": use["all"], "mine": use["mine"], "running": use["jobs"], "my_running": use["my_jobs"],
        })
    qos = [dict(zip(("name", "group", "per_user", "jobs_per_user", "max_wall"), (f + [""] * 5)[:5])) for f in rows("qos") if f and f[0]]

    by_part: dict[str, dict] = {}
    for f in rows("nodes"):
        if len(f) < 6:
            continue
        part = f[0].rstrip("*")
        state = f[2].rstrip("*~#!%$@^-").lower()
        p = by_part.setdefault(part, {"states": {}, "mem_free_mb": 0, "mem_total_mb": 0, "cpus_idle": 0})
        p["states"][state] = p["states"].get(state, 0) + 1
        usable = state in ("idle", "mixed", "mix", "allocated", "alloc", "completing")
        mem = float(f[4]) if f[4].isdigit() else 0
        free = float(f[5]) if f[5].isdigit() else 0
        if usable:
            p["mem_total_mb"] += mem
            if state in ("idle", "mixed", "mix"):
                p["mem_free_mb"] += free
                p["cpus_idle"] += _cpus(f[3])["idle"]
    partitions = []
    for f in rows("partitions"):
        if len(f) < 7:
            continue
        name = f[0].rstrip("*")
        nodes = by_part.get(name, {})
        partitions.append({
            "name": name, "default": f[0].endswith("*"), "up": f[1] == "up", "time_limit": f[2],
            "nodes": int(f[3]) if f[3].isdigit() else 0, "cpus": _cpus(f[4]),
            "mem_per_node_mb": int(f[5].rstrip("+")) if f[5].rstrip("+").isdigit() else 0,
            "gres": "" if f[6] == "(null)" else f[6],
            "node_states": nodes.get("states", {}),
            "cpus_ready": nodes.get("cpus_idle", 0),
            "mem_free_mb": nodes.get("mem_free_mb", 0), "mem_total_mb": nodes.get("mem_total_mb", 0),
        })
    return {
        "host": host,
        "on_login_node": "login" in host.split(".")[0].lower() or not host,
        "jobs": jobs,
        "interactive": [j for j in jobs if j["interactive"] and j["state"] == "RUNNING"],
        "accounts": accounts,
        "qos": qos,
        "partitions": partitions,
        "complete": "end" in sections,
    }


@router.get("/cluster_resources")
def cluster_resources(current_user: dict = Depends(get_current_user)):
    """The user's SLURM picture: accounts with their limits and what is in use, QOS limits, the
    user's jobs (and any interactive session, with its node), and each partition's free CPUs and
    memory. Reused for 20 seconds (read_cache)."""
    conn = _build_connection(current_user)

    def ask():
        _, out = _run_remote_check(conn, CLUSTER_SCRIPT)
        result = parse_cluster(out, current_user["cluster_username"])
        if not result["complete"] and not result["partitions"]:
            raise HTTPException(status_code=502, detail="The cluster did not answer the SLURM questions; is SLURM there?")
        return result

    return read_cache.get((current_user["user_id"], "cluster"), ask, fresh=20, stale=90)


@router.get("/status")
def ssh_status(current_user: dict = Depends(get_current_user)):
    """Checks whether the server can reach the user's cluster over SSH.

    Always returns 200, with the reason for a failure and whether the user can fix it.
    """
    try:
        conn = _build_connection(current_user)
        ssh = conn.connect()
        # Not closed: the client is pooled and shared across requests.
        return {"connected": True, "host": current_user["cluster_host"]}
    except HTTPException as exc:
        # Undecryptable stored key (BSP_ENCRYPTION_KEY regenerated): re-register.
        detail = str(getattr(exc, "detail", exc))
        undecryptable = "decrypt" in detail.lower()
        LOGGER.warning("SSH status check failed for user %s: %s",
                       current_user["username"], detail)
        return {
            "connected": False,
            "host": current_user["cluster_host"],
            "reason": "key_undecryptable" if undecryptable else "error",
            "detail": (
                "Your stored SSH key cannot be decrypted, because the server's "
                "encryption key changed after this account was created. The key "
                "cannot be recovered — please register a new account to continue."
                if undecryptable else detail
            ),
            "action": "re-register" if undecryptable else None,
        }
    except Exception as exc:
        LOGGER.warning("SSH status check failed for user %s: %s",
                       current_user["username"], exc)
        return {
            "connected": False,
            "host": current_user["cluster_host"],
            "reason": "unreachable",
            "detail": f"Could not reach {current_user['cluster_host']}: {exc}",
            "action": None,
        }


@router.get("/config")
def get_config(current_user: dict = Depends(get_current_user)):
    """Read the user's ~/.config/bioinformatics-tools/config.yaml from their cluster via SFTP."""
    conn = _build_connection(current_user)
    path = _config_path(current_user["home_dir"])
    try:
        return read_cache.get((current_user["user_id"], "config"), lambda: ssh_sftp.read_remote_yaml(path, connection=conn))
    except Exception as exc:
        LOGGER.error("Failed to read remote config for %s: %s", current_user["username"], exc)
        raise HTTPException(status_code=500, detail=f"Failed to read remote config: {exc}")


@router.put("/config")
def save_config(config: dict, current_user: dict = Depends(get_current_user)):
    """Write a config dict back to the user's cluster as YAML via SFTP."""
    conn = _build_connection(current_user)
    path = _config_path(current_user["home_dir"])
    try:
        ssh_sftp.write_remote_yaml(path, config, connection=conn)
        return {"success": True}
    except Exception as exc:
        LOGGER.error("Failed to write remote config for %s: %s", current_user["username"], exc)
        raise HTTPException(status_code=500, detail=f"Failed to write remote config: {exc}")


@router.post("/config/create-default")
def create_default_config(current_user: dict = Depends(get_current_user)):
    """Create a default config file with all system defaults populated."""
    conn = _build_connection(current_user)
    path = _config_path(current_user["home_dir"])

    default_config = _build_default_config_payload(current_user.get("cluster_username"))
    default_config_text = _build_default_config_text(default_config)

    try:
        ssh_sftp.write_remote_text_file(path, default_config_text, connection=conn)
        LOGGER.info("Created default config for user %s at %s", current_user["username"], path)
        return {"success": True, "config": default_config}
    except Exception as exc:
        LOGGER.error("Failed to create default config for %s: %s", current_user["username"], exc)
        raise HTTPException(status_code=500, detail=f"Failed to create default config: {exc}")


@router.post("/test-path-writable")
def test_path_writable(path_data: dict, current_user: dict = Depends(get_current_user)):
    """Test if a path on the cluster is writable by attempting to create parent directories and a test file."""
    conn = _build_connection(current_user)
    test_path = path_data.get("path", "").strip()

    if not test_path:
        raise HTTPException(status_code=400, detail="Path is required")

    try:
        ssh = conn.connect()

        # Expand ~ to actual home directory
        if test_path.startswith("~"):
            test_path = test_path.replace("~", current_user["home_dir"], 1)

        # Get the directory (remove filename if present)
        import posixpath
        test_dir = posixpath.dirname(test_path)

        # Try to create the directory structure
        _, stdout, stderr = ssh.exec_command(f'mkdir -p "{test_dir}" 2>&1 && echo "DIR_OK"')
        output = stdout.read().decode().strip()

        if "DIR_OK" not in output:
            pass  # pooled client: closing it would break concurrent requests (see SSHConnection pool)
            return {
                "writable": False,
                "error": f"Cannot create directory: {test_dir}",
                "details": output
            }

        # Try to write a test file
        test_file = f"{test_path}.write_test"
        _, stdout, stderr = ssh.exec_command(f'touch "{test_file}" 2>&1 && rm -f "{test_file}" 2>&1 && echo "WRITE_OK"')
        output = stdout.read().decode().strip()

        pass  # pooled client: closing it would break concurrent requests (see SSHConnection pool)

        if "WRITE_OK" in output:
            return {"writable": True}
        else:
            return {
                "writable": False,
                "error": f"Path is not writable: {test_path}",
                "details": output
            }

    except Exception as exc:
        LOGGER.error("Failed to test path writability for %s: %s", current_user["username"], exc)
        return {
            "writable": False,
            "error": f"Failed to test path: {str(exc)}"
        }


@router.post("/run_slurm")
def run_slurm(content: SlurmSend, current_user: dict = Depends(get_current_user)):
    """Submit a SLURM job and return the job ID immediately."""
    conn = _build_connection(current_user)
    job_id = ssh_slurm.submit_slurm_job(script_content=content.script, connection=conn)
    return {"success": True, "job_id": job_id, "message": "Job submitted successfully"}


@router.post("/run_ssh")
def run_ssh(content: SlurmSend, current_user: dict = Depends(get_current_user)):
    """Execute an SSH command and return output."""
    LOGGER.info('Running run_ssh for user %s', current_user["username"])
    conn = _build_connection(current_user)
    std_txt = ssh_slurm.submit_ssh_job(cmd=content.script, connection=conn)
    return {"success": True, "std_txt": std_txt, "message": "Job submitted successfully"}


def _check_genome_path_exists(genome_path: str, workflow: str, conn) -> None:
    """Raises HTTPException(400) if genome_path is missing on the cluster, or a
    batch-input folder holds no recognised genome file. Used for new runs and relaunches."""
    wf_key = WORKFLOWS.get(workflow)
    supports_batch_input = bool(wf_key and wf_key.supports_batch_input)
    try:
        if supports_batch_input:
            attr = ssh_sftp.check_remote_path_kind(genome_path, conn)
            if attr == 'directory':
                entries = ssh_sftp.list_remote_dir(genome_path, conn)
                has_genome_file = any(
                    e['type'] == 'file' and e['name'].lower().endswith(GENOME_EXTENSIONS)
                    for e in entries
                )
                if not has_genome_file:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No recognized genome files (e.g. .fasta, .fa, .fna) found in folder: '{genome_path}'",
                    )
        else:
            ssh_sftp.check_remote_file(genome_path, conn)
    except FileNotFoundError:
        raise HTTPException(
            status_code=400,
            detail=f"Path not found on the cluster: '{genome_path}'. "
                   "Make sure the path is on the cluster, not on your local machine.",
        )
    except IsADirectoryError:
        raise HTTPException(
            status_code=400,
            detail=f"Path points to a directory, not a file: '{genome_path}'",
        )
    except HTTPException:
        raise
    except Exception as exc:
        LOGGER.warning("File pre-check failed for %s: %s", genome_path, exc)
        raise HTTPException(
            status_code=400,
            detail=f"Could not verify path on cluster: {exc}",
        )


#: Beside each job's results: the exact command it ran, to read, copy or run again.
COMMAND_FILE = "run-command.sh"


def _dane_wf_command(workflow: str, genome_path: str, output_dir: str, selected_tools_csv: str | None, *,
                     resume: bool = False, full_operon_map: bool = False, license_env: str = "") -> str:
    """The dane_wf command line a job runs (caragols dispatches do_margie_sb as "margie sb")."""
    dispatch_tokens = workflow.replace('_', ' ')
    selected_tools_arg = f" {workflow}.selected_tools: {selected_tools_csv}" if selected_tools_csv else ""
    # Opt-in full operon atlas; the smk runs it after the report figures.
    full_operon_map_arg = " run_full_operon_map: true" if full_operon_map else ""
    resume_arg = f" {workflow}.resume: true" if resume else ""
    return (
        f"{license_env}~/bioinformatics-tools/.venv/bin/dane_wf {dispatch_tokens}"
        f" input: {genome_path} output_dir: {output_dir}{selected_tools_arg}{full_operon_map_arg}{resume_arg}"
    )


def _command_script(job_id: str, workflow: str, genome_path: str, output_dir: str, selected_tools_csv: str | None, *,
                    resume: bool = False, full_operon_map: bool = False, env: list[tuple[str, str]] | None = None,
                    note: str = "") -> str:
    """The same command as _dane_wf_command, set out to run again from a shell on the cluster:
    from the code's folder, one argument to a line."""
    head = [
        "#!/usr/bin/env bash",
        f"# The command MARGIE ran for job {job_id}, set out to run again from a shell on the cluster.",
        "# MARGIE started it inside a SLURM job; run from here it works the same way, sending each step to SLURM.",
    ]
    if note:
        head.append(f"# {note}")
    parts = [f"{k}='{v}'" for k, v in (env or [])]
    parts.append(f".venv/bin/dane_wf {workflow.replace('_', ' ')}")
    parts += [f"  input: {genome_path}", f"  output_dir: {output_dir}"]
    if selected_tools_csv:
        parts.append(f"  {workflow}.selected_tools: {selected_tools_csv}")
    if full_operon_map:
        parts.append("  run_full_operon_map: true")
    if resume:
        parts.append(f"  {workflow}.resume: true")
    body = [f"{line} \\" for line in parts[:-1]] + [parts[-1]]
    return "\n".join(head + ["cd ~/bioinformatics-tools", 'export PATH="$HOME/.local/bin:$PATH"'] + body) + "\n"


def _save_command(conn, output_dir: str, job_id: str, text: str) -> None:
    """Writes the job's command script into its output folder as run-command.sh. Never stops a launch."""
    folder = output_dir.rstrip('/')
    try:
        ssh = conn.connect()
        _, out, _ = ssh.exec_command(f"mkdir -p {shlex.quote(folder)}")
        out.channel.recv_exit_status()
        ssh_sftp.write_remote_text_file(f"{folder}/{COMMAND_FILE}", text, connection=conn)
    except Exception as exc:
        LOGGER.warning("Could not save the command for job %s: %s", job_id, exc)


def _read_small_remote(path: str, conn, limit: int = 65536) -> str:
    """Reads a small text file on the cluster. Raises OSError when it is not there."""
    sftp = conn.connect().open_sftp()
    try:
        with sftp.open(path, 'r') as f:
            return f.read(limit).decode(errors='replace')
    finally:
        sftp.close()


def _claim_run_folder(conn, wanted: str) -> str:
    """Makes the run's own folder: `wanted`, or wanted-2, wanted-3, ... when a run already has it.

    Run folders are named to the minute, so a run started in the same minute as another (a
    cancelled one started again) would otherwise share its folder and reuse its outputs. mkdir
    either makes the folder or fails, so two launches never claim the same one. Returns `wanted`
    when the cluster gives no answer; dane_wf then makes it."""
    q = shlex.quote(wanted)
    try:
        _, out = _run_remote_check(
            conn,
            f'mkdir -p -- "$(dirname -- {q})" 2>/dev/null; p={q}; n=1; '
            f'until mkdir -- "$p" 2>/dev/null; do n=$((n+1)); [ "$n" -gt 99 ] && exit 1; p={q}-$n; done; printf "%s\\n" "$p"',
        )
    except Exception as exc:  # never blocks the launch
        LOGGER.warning("Could not claim run folder %s: %s", wanted, exc)
        return wanted
    claimed = out.strip().splitlines()[-1] if out.strip() else ""
    return claimed if claimed == wanted or claimed.startswith(f"{wanted}-") else wanted


def _launch_job(
    *, genome_path: str, workflow: str, base_output_dir: str,
    selected_tools: list[str] | None, current_user: dict, conn, main_db: str | None,
    slurm_account: str | None = None,
    slurm_partition: str | None = None,
    slurm_walltime: str | None = None,
    relaunched_from: str | None = None,
    copy_from_work_dir: str | None = None,
    run_full_operon_map: bool = False,
) -> dict:
    """Generates job_id and output_dir, optionally copies a previous run's output
    forward (Resume, with margie_sb.resume: true), records the job and submits dane_wf.

    Pre-flight validation is left to the callers.
    """
    selected_tools_csv = ",".join(selected_tools) if selected_tools is not None else None

    job_id = str(uuid.uuid4())
    timestamp = datetime.now().strftime('%Y-%m-%d-%H%M')
    output_dir = _claim_run_folder(conn, f"{base_output_dir.rstrip('/')}/{timestamp}")

    if copy_from_work_dir:
        ssh_sftp.copy_remote_directory(copy_from_work_dir, output_dir, connection=conn)
        try:
            ssh_sftp.rewrite_path_references(output_dir, copy_from_work_dir, output_dir, connection=conn)
        except Exception as exc:
            # Cosmetic provenance cleanup; a failure never blocks the launch.
            LOGGER.warning("Could not rewrite stale path references for resumed job: %s", exc)

    job_store.create(
        job_id, genome_path, user_id=current_user["user_id"],
        workflow=workflow, output_dir=output_dir,
        selected_tools=selected_tools_csv, relaunched_from=relaunched_from,
        persist_owner_username=current_user["username"],
        persist_owner_cluster_username=current_user["cluster_username"],
        persist_db_path=main_db, persist_connection=conn,
    )
    job_store.update(job_id, work_dir=output_dir)

    # Passes the verified acceptance so the CLI licence gate does not re-prompt.
    _lic_ent = licensing.get_entitlement(current_user["username"])
    _lic_csv = ",".join(_lic_ent.get("licensed_tools") or [])
    license_pairs = [
        ("MARGIE_LICENSE_ACCEPTED", str(licensing.load_terms()['version'])),
        ("MARGIE_USAGE_TYPE", str(_lic_ent.get('usage_type') or '')),
        ("MARGIE_LICENSED_TOOLS", _lic_csv),
    ]
    license_env = "".join(f"{k}='{v}' " for k, v in license_pairs)
    # Uses the editable install, not `uvx --from`, which is slow and may be stale.
    if workflow == 'margie_sb':
        try:
            # A run uses the account's own clone of the repository, so one not set up yet
            # (or still an older checkout) is set up first.
            if sync_remote_dane_wf(conn) in ('unprovisioned', 'not-a-git-checkout', 'other-repo'):
                provision_in_background(conn)
                wait_provisioned(conn)
        except Exception as exc:
            LOGGER.warning("dane_wf version-sync check raised unexpectedly, ignoring: %s", exc)

    command = _dane_wf_command(
        workflow, genome_path, output_dir, selected_tools_csv,
        resume=bool(copy_from_work_dir), full_operon_map=run_full_operon_map, license_env=license_env,
    )
    _save_command(conn, output_dir, job_id, _command_script(
        job_id, workflow, genome_path, output_dir, selected_tools_csv,
        resume=bool(copy_from_work_dir), full_operon_map=run_full_operon_map, env=license_pairs,
    ))
    job_runner.submit_job(
        job_id,
        command,
        connection=conn,
        driver_account=slurm_account,
        driver_partition=slurm_partition,
        driver_time=slurm_walltime,
    )

    return {"success": True, "job_id": job_id, "output_dir": output_dir, "message": "Job submitted successfully"}


# ---- Per-user databases on scratch and their depot backups ----

STORES_NOT_READY = "stores-not-ready"


def _stores_user_config(current_user: dict, conn) -> dict:
    try:
        return ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn)
    except Exception:
        raise HTTPException(status_code=400, detail="Configuration file not found. Create your configuration first.")


def _stores_call(fn, *args):
    try:
        return fn(*args)
    except user_stores.StoreError as exc:
        raise HTTPException(status_code=exc.status, detail=str(exc))


def _stores_settle(current_user: dict, conn, user_config: dict) -> None:
    """Records a finished copy, and saves the config if it now points elsewhere."""
    user = current_user["cluster_username"]
    if _stores_call(user_stores.apply_finished, conn, user_config, user):
        ssh_sftp.write_remote_yaml(_config_path(current_user["home_dir"]), user_config, connection=conn)
        read_cache.drop(current_user["user_id"], "config", "jobs")


def _margie_sb_stores_ready(current_user: dict, conn, user_config: dict) -> str:
    """Returns the job database for a margie_sb run once every store is on scratch.

    Raises 409 "stores-not-ready..." when they are not, or while a copy is under way."""
    user = current_user["cluster_username"]
    _stores_settle(current_user, conn, user_config)
    st = _stores_call(user_stores.status, conn, user_config, user)
    op = st.get("op") or {}
    if op.get("state") == "running":
        raise HTTPException(status_code=409, detail=f"{STORES_NOT_READY}: your databases are being copied ({op.get('label', '')}); the run can start when that finishes.")
    if not st["ready"]:
        raise HTTPException(status_code=409, detail=f"{STORES_NOT_READY}: your databases are not on scratch yet. They are copied there once, before your first run.")
    if user_stores.point_config(conn, user_config, user):
        ssh_sftp.write_remote_yaml(_config_path(current_user["home_dir"]), user_config, connection=conn)
    return user_config["main_database"]


def _active_run(current_user: dict, conn, user_config: dict) -> bool:
    """Returns whether any of this user's runs is still going (checked live)."""
    main_db = user_config.get("main_database")
    if not main_db:
        return False
    try:
        rows, _ = job_history_client.list_jobs_and_count(
            conn, main_db, limit=20, offset=0,
            owner_username=current_user["username"],
            owner_cluster_username=current_user["cluster_username"],
        )
        _reconcile_running(conn, main_db, rows, current_user["cluster_username"])
    except Exception as exc:
        LOGGER.warning("Could not check for running jobs before a backup: %s", exc)
        return False
    return any(str(r.get("status", "")).lower() == "running" for r in rows)


@router.get("/stores")
def get_stores(current_user: dict = Depends(get_current_user)):
    """Returns each database's location and version, latest depot backup, and the
    current or last copy with its percentage and log tail."""
    conn = _build_connection(current_user)

    def ask():
        user_config = _stores_user_config(current_user, conn)
        _stores_settle(current_user, conn, user_config)
        return _stores_call(user_stores.status, conn, user_config, current_user["cluster_username"])

    return read_cache.get((current_user["user_id"], "stores"), ask)


@router.get("/stores/progress")
def get_stores_progress(current_user: dict = Depends(get_current_user)):
    """Returns only the copy under way (percent, log tail); cheap enough to poll often."""
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    return {"op": _stores_call(user_stores.progress, conn, user_config)}


@router.get("/stores/{store_id}/backup-check")
def check_store_backup(store_id: str, current_user: dict = Depends(get_current_user)):
    """Returns what backing up this database would take and whether depot has room."""
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    return _stores_call(user_stores.backup_check, conn, user_config, current_user["cluster_username"], store_id)


@router.post("/stores/setup")
def setup_stores(body: dict | None = None, current_user: dict = Depends(get_current_user)):
    """Copies the databases missing on scratch, from the user's newest depot
    backup where there is one, else from the base copies.

    body {"fresh": true, "build_password": ...}: every database starts empty instead, as its
    next version in the working folder (a new lab's own); this needs the build password."""
    fresh = bool((body or {}).get("fresh"))
    if fresh:
        build_gate.check((body or {}).get("build_password"))
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    _stores_settle(current_user, conn, user_config)
    return _stores_call(user_stores.start_setup, conn, user_config, current_user["cluster_username"], fresh)


# ---- Genome viewer refresh (rebuilt beside the old file, then moved over) ----

@router.post("/genome-viewer/refresh")
def refresh_genome_viewer(body: dict | None = None, current_user: dict = Depends(get_current_user)):
    """Rebuilds a genome's viewer. body: {folder}; returns {viewer}: the map's path."""
    folder = str((body or {}).get("folder") or "").strip()
    if not folder:
        raise HTTPException(status_code=400, detail="Which genome folder?")
    folder = _resolve_browse_path(folder, current_user)
    conn = _build_connection(current_user)
    q = shlex.quote
    home = current_user["home_dir"]
    repo = f"{home}/bioinformatics-tools"
    script = f"{repo}/bioinformatics_tools/workflow_tools/viz/gen_genome_viewer.py"
    cmd = (
        f"cd {q(folder)} || exit 3; "
        "T=FINAL_ANNOTATION_WITH_CONFIDENCE.tsv; [ -f \"$T\" ] || T=scoring/$T; [ -f \"$T\" ] || exit 4; "
        "V=FINAL_GENOME_VIEWER.html; [ -f \"$V\" ] || { [ -f scoring/$V ] && V=scoring/$V; }; "
        f"PY=${{MARGIE_PYTHON:-{q(repo)}/.venv/bin/python}}; "
        f"\"$PY\" {q(script)} \"$T\" \"$V.new\" > .viewer-refresh.log 2>&1 && mv -f \"$V.new\" \"$V\" && echo \"$V\""
    )
    code, out = user_stores._run(conn, cmd, timeout=600)
    if code == 3:
        raise HTTPException(status_code=404, detail=f"No folder {folder} on the cluster.")
    if code == 4:
        raise HTTPException(status_code=404, detail="This genome has no report table yet, so there is no map to make.")
    if code != 0:
        _, log = user_stores._run(conn, f"tail -n 15 {q(folder)}/.viewer-refresh.log 2>/dev/null")
        raise HTTPException(status_code=500, detail=f"The map could not be made again: {(log or out).strip()[-600:]}")
    viewer = out.strip().splitlines()[-1] if out.strip() else "FINAL_GENOME_VIEWER.html"
    return {"viewer": f"{folder}/{viewer}"}


# ---- AI keys: ~/.config/margie on the cluster, JSON or provider=key lines ----

_AI_KEYS_DIR = ".config/margie"
_AI_KEYS_FILE = "ai-keys.json"
_AI_PROVIDERS = {"anthropic", "openai", "gemini", "custom"}


@router.get("/ai-keys")
def get_ai_keys(current_user: dict = Depends(get_current_user)):
    """Returns the keys kept on the cluster, by provider ({} when none)."""
    conn = _build_connection(current_user)
    path = f"{current_user['home_dir']}/{_AI_KEYS_DIR}/{_AI_KEYS_FILE}"
    code, out = user_stores._run(conn, f"cat {shlex.quote(path)} 2>/dev/null")
    text = out if code == 0 else ""
    keys: dict = {}
    try:
        data = json.loads(text) if text.strip() else {}
        keys = (data.get("keys") if isinstance(data, dict) else None) or {}
    except json.JSONDecodeError:
        # Hand-written form: one provider=key line each.
        for line in text.splitlines():
            name, sep, value = line.strip().partition("=")
            if sep and not line.lstrip().startswith("#"):
                keys[name.strip().lower()] = value.strip().strip("'\"")
    return {"keys": {k: v for k, v in keys.items() if k in _AI_PROVIDERS and isinstance(v, str) and v}}


@router.put("/ai-keys")
def put_ai_keys(body: dict | None = None, current_user: dict = Depends(get_current_user)):
    """Replaces the keys kept on the cluster. body: {keys: {provider: key}}."""
    keys = (body or {}).get("keys") or {}
    if not isinstance(keys, dict):
        raise HTTPException(status_code=400, detail="Expected {keys: {provider: key}}.")
    clean = {k: str(v).strip() for k, v in keys.items() if k in _AI_PROVIDERS and isinstance(v, str) and v.strip()}
    conn = _build_connection(current_user)
    home = current_user["home_dir"]
    folder = f"{home}/{_AI_KEYS_DIR}"
    code, out = user_stores._run(conn, f"mkdir -p {shlex.quote(folder)} && chmod 700 {shlex.quote(folder)}")
    if code != 0:
        raise HTTPException(status_code=500, detail=f"Could not make {folder}: {out.strip()}")
    path = f"{folder}/{_AI_KEYS_FILE}"
    ssh_sftp.write_remote_text_file(path, json.dumps({"keys": clean}) + "\n", connection=conn)
    user_stores._run(conn, f"chmod 600 {shlex.quote(path)}")
    return {"saved": sorted(clean)}


@router.delete("/ai-keys")
def delete_ai_keys(current_user: dict = Depends(get_current_user)):
    """Deletes the keys kept on the cluster."""
    conn = _build_connection(current_user)
    path = f"{current_user['home_dir']}/{_AI_KEYS_DIR}/{_AI_KEYS_FILE}"
    user_stores._run(conn, f"rm -f {shlex.quote(path)}")
    return {"ok": True}


@router.post("/stores/run-here")
def run_store_copy_here(current_user: dict = Depends(get_current_user)):
    """Runs the queued SLURM copy job (databases, backups, tool setup) on the login node instead."""
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    return {"op": _stores_call(user_stores.run_here, conn, user_config)}


@router.post("/stores/stop")
def stop_store_copy(current_user: dict = Depends(get_current_user)):
    """Stops the copy under way (databases, backups, tool setup), queued or running; see user_stores.stop."""
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    return {"op": _stores_call(user_stores.stop, conn, user_config)}


@router.post("/stores/{store_id}/backup")
def backup_store(store_id: str, current_user: dict = Depends(get_current_user)):
    """Copies one database's working version to depot; the working copy moves to
    the next version. Refused while a run is writing to it."""
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    _stores_settle(current_user, conn, user_config)
    if _active_run(current_user, conn, user_config):
        raise HTTPException(status_code=409, detail="A run is still going; back up once it has finished, so the copy is not taken half-written.")
    return _stores_call(user_stores.start_backup, conn, user_config, current_user["cluster_username"], store_id)


# ---- Tool containers and reference databases (share the stores copy job) ----

@router.get("/assets")
def get_assets(current_user: dict = Depends(get_current_user)):
    """Returns each container and reference database: where the run looks for it,
    whether it is there, and where it could be set up from."""
    conn = _build_connection(current_user)

    def ask():
        user_config = _stores_user_config(current_user, conn)
        _stores_settle(current_user, conn, user_config)
        return _stores_call(tool_assets.status, conn, user_config, current_user["home_dir"])

    return read_cache.get((current_user["user_id"], "assets"), ask)


@router.get("/assets/plan")
def get_assets_plan(builds: bool = True, optional: bool = False, accept: str = "",
                    sif_to: str = "", db_to: str = "",
                    current_user: dict = Depends(get_current_user)):
    """Returns the setup plan: items, copy sizes and room on scratch.
    accept=* also lists the licence-gated builds, marked."""
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    accepted = set(tool_assets.GATED) if accept == "*" else {t.strip() for t in accept.split(",") if t.strip()}
    return _stores_call(tool_assets.plan, conn, user_config, current_user["home_dir"], builds, accepted, optional,
                        sif_to or None, db_to or None)


@router.post("/assets/setup")
def setup_assets(body: dict | None = None, current_user: dict = Depends(get_current_user)):
    """Sets up the missing containers and databases.

    body: {builds, optional, accept: [tool, ...], sif_to, db_to}; accept names the
    licence-gated tools the user accepted, sif_to / db_to the chosen folders."""
    body = body or {}
    # Copying is open to everyone; building needs the build password (api/build_gate.py).
    if bool(body.get("builds", True)):
        build_gate.check(body.get("build_password"))
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    _stores_settle(current_user, conn, user_config)
    accepted = {str(t) for t in (body.get("accept") or [])}
    return _stores_call(tool_assets.start, conn, user_config, current_user["home_dir"],
                        bool(body.get("builds", True)), accepted, bool(body.get("optional", False)),
                        str(body.get("sif_to") or "") or None, str(body.get("db_to") or "") or None)


@router.get("/assets/build-recipes")
def get_build_recipes(current_user: dict = Depends(get_current_user)):
    """Returns where margie-build is on the cluster (shared copy or user clone), if anywhere."""
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    return _stores_call(tool_assets.build_recipes, conn, user_config, current_user["home_dir"])


@router.post("/assets/build-unlock")
def build_unlock(body: dict | None = None, current_user: dict = Depends(get_current_user)):
    """Checks the build password, so the page can open the building controls. body: {password}."""
    build_gate.check((body or {}).get("password"))
    return {"ok": True}


@router.post("/assets/build-recipes")
def clone_build_recipes(body: dict | None = None, current_user: dict = Depends(get_current_user)):
    """Clones margie-build into the user's home and points margie_sb.build_repo at it.
    body: {url, build_password}; cloning is part of building, so it needs the password."""
    build_gate.check((body or {}).get("build_password"))
    conn = _build_connection(current_user)
    user_config = _stores_user_config(current_user, conn)
    found = _stores_call(tool_assets.clone_recipes, conn, user_config, current_user["home_dir"], (body or {}).get("url"))
    ssh_sftp.write_remote_yaml(_config_path(current_user["home_dir"]), user_config, connection=conn)
    return found


def _check_margie_sb_assets(genome_data: GenomeSend, genome_path: str, user_config: dict, conn, current_user: dict) -> None:
    """Refuses a run whose containers or databases are missing. An unreadable
    folder or a failed check never stops a run."""
    all_keys = {tool['key'] for tool in MARGIE_SB_PHASED_TOOLS}
    if genome_data.selected_tools is not None:
        tools = set(genome_data.selected_tools)
    else:
        ent = licensing.get_entitlement(current_user["username"])
        tools = all_keys - licensing.disabled_tool_ids(ent.get("usage_type"), ent.get("licensed_tools"))
    names = genome_data.genomes
    if names is None:
        try:
            if ssh_sftp.check_remote_path_kind(genome_path, conn) == 'directory':
                names = [e['name'] for e in ssh_sftp.list_remote_dir(genome_path, conn)
                         if e.get('type') == 'file' and e['name'].lower().endswith(GENOME_EXTENSIONS)]
            else:
                names = [posixpath.basename(genome_path)]
        except Exception:
            names = None
    try:
        missing = tool_assets.missing_for_run(conn, user_config, current_user["home_dir"], tools, names)
    except Exception as exc:
        LOGGER.warning("Could not check the containers and databases before a run: %s", exc)
        return
    if missing:
        raise HTTPException(status_code=409, detail=tool_assets.describe_missing(missing))


def _require_current_terms(current_user: dict) -> None:
    """Server-side licence gate: nothing runs, resumes or restarts without acceptance of the current terms."""
    if not licensing.has_accepted_current_terms(current_user["username"]):
        raise HTTPException(
            status_code=403,
            detail="You must accept the current MARGIE licensing terms before running, resuming or restarting an analysis.",
        )


@router.post("/run_workflow")
def run_workflow(genome_data: GenomeSend, current_user: dict = Depends(get_current_user)):
    """Submit a genome analysis workflow by name."""
    _require_current_terms(current_user)

    available_workflows = _get_available_workflows()
    allowed_ids = {wf["id"] for wf in available_workflows}

    if genome_data.workflow not in allowed_ids:
        raise HTTPException(status_code=400, detail=f"Unknown workflow '{genome_data.workflow}'. Available: {sorted(allowed_ids)}")

    if genome_data.workflow in STUB_WORKFLOWS:
        raise HTTPException(status_code=501, detail=f"Workflow '{genome_data.workflow}' is not yet implemented. Check back soon!")

    conn = _build_connection(current_user)

    # Pre-flight: validate required config values are set
    config_path = _config_path(current_user["home_dir"])
    try:
        user_config = ssh_sftp.read_remote_yaml(config_path, connection=conn)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Configuration file not found. Please create a configuration in your Profile settings first."
        )

    # Validate required fields
    missing_fields = []

    # Check main_database
    main_db = user_config.get('main_database')
    if not main_db or str(main_db).strip() == '':
        missing_fields.append('main_database')

    # Check compute.cluster_default.account
    account = user_config.get('compute', {}).get('cluster_default', {}).get('account')
    if not account or str(account).strip() == '':
        missing_fields.append('compute.cluster_default.account (SLURM account)')
    slurm_account = str(account).strip() if account else None
    partition = user_config.get('compute', {}).get('cluster_default', {}).get('partition')
    slurm_partition = str(partition).strip() if partition else None
    slurm_walltime = _cluster_default(user_config, 'driver_walltime')

    if missing_fields:
        raise HTTPException(
            status_code=400,
            detail=f"Required configuration missing: {', '.join(missing_fields)}. "
                   "Please configure these in your Profile settings before running workflows."
        )

    if genome_data.workflow == 'margie_sb':
        # Store copies must be on scratch; never starts during a copy.
        main_db = _margie_sb_stores_ready(current_user, conn, user_config)
        _validate_margie_sb_shared_paths(user_config, conn, current_user["home_dir"])
    else:
        # Promotes a shared template DB to a per-user DB (no shared SQLite file).
        main_db = _resolve_effective_main_db(current_user, conn, user_config, persist=True)

    # Falls back to the config's input_path / output_path.
    genome_path = genome_data.genome_path or user_config.get(genome_data.workflow, {}).get('input_path')
    if not genome_path or str(genome_path).strip() == '':
        raise HTTPException(
            status_code=400,
            detail="No genome file or folder specified, and no input_path default is configured. "
                   "Set one in your Profile settings or pass genome_path explicitly.",
        )

    _check_genome_path_exists(genome_path, genome_data.workflow, conn)

    # A folder subset is staged into its own folder; an empty list is rejected.
    if genome_data.genomes is not None:
        if not genome_data.genomes:
            raise HTTPException(status_code=400, detail="No genomes were selected.")
        if ssh_sftp.check_remote_path_kind(genome_path, conn) != 'directory':
            raise HTTPException(
                status_code=400,
                detail="Choosing genomes needs a folder to choose from; this run points at a single file.",
            )
        try:
            genome_path = ssh_sftp.stage_selected_genomes(
                genome_path, genome_data.genomes, conn, label=current_user["username"],
            )
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        except RuntimeError as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc
        LOGGER.info(
            "Running %d selected genome(s) from a staged folder: %s",
            len(genome_data.genomes), genome_path,
        )

    # Validates tool keys; [] means "run nothing", None "run everything".
    if genome_data.selected_tools is not None:
        if not genome_data.selected_tools:
            raise HTTPException(
                status_code=400,
                detail="selected_tools was empty -- select at least one tool/phase to run, "
                       "or omit the field entirely to run everything.",
            )
        valid_tool_keys = {tool['key'] for tool in MARGIE_SB_PHASED_TOOLS}
        unknown = set(genome_data.selected_tools) - valid_tool_keys
        if unknown:
            raise HTTPException(
                status_code=400,
                detail=f"Unknown tool key(s) in selected_tools: {sorted(unknown)}. "
                       f"Available: {sorted(valid_tool_keys)}",
            )
        # Refuses tools the user is not licensed for (mirrors the CLI gate).
        _ent = licensing.get_entitlement(current_user["username"])
        _disabled = licensing.disabled_tool_ids(
            _ent.get("usage_type"), _ent.get("licensed_tools")
        ) & valid_tool_keys
        _blocked = set(genome_data.selected_tools) & _disabled
        if _blocked:
            raise HTTPException(
                status_code=403,
                detail=f"These tools require a license you have not recorded: {sorted(_blocked)}. "
                       "Update your usage type / licensed tools when accepting the terms "
                       "(Profile), or remove them from your selection.",
            )

    if genome_data.workflow == 'margie_sb':
        _check_margie_sb_assets(genome_data, genome_path, user_config, conn, current_user)

    base_dir = (genome_data.output_dir or user_config.get(genome_data.workflow, {}).get('output_path') or current_user['home_dir']).rstrip('/')

    saved_full_operon_map = bool(
        user_config.get(genome_data.workflow, {}).get('run_full_operon_map', False))
    effective_full_operon_map = bool(genome_data.run_full_operon_map) or saved_full_operon_map

    return _launch_job(
        genome_path=genome_path, workflow=genome_data.workflow,
        base_output_dir=base_dir, selected_tools=genome_data.selected_tools,
        current_user=current_user, conn=conn, main_db=main_db,
        slurm_account=slurm_account, slurm_partition=slurm_partition,
        slurm_walltime=slurm_walltime,
        run_full_operon_map=effective_full_operon_map,
    )


def _job_from_history_row(row: dict) -> dict:
    """Shapes a persisted api_jobs row like a live job_store entry.

    logs/slurm_jobs/containers exist only if job_store.finalize() ran; sub_jobs,
    report and progress are never persisted.
    """
    return {
        "job_id": row["job_id"],
        "owner_username": row.get("owner_username"),
        "owner_cluster_username": row.get("owner_cluster_username"),
        "status": row["status"],
        "phase": row.get("phase"),
        "genome_path": row.get("genome_path"),
        "workflow": row.get("workflow"),
        "work_dir": row.get("work_dir"),
        "selected_tools": row.get("selected_tools"),
        "relaunched_from": row.get("relaunched_from"),
        "start_time": row.get("created_at"),
        # A finished row's updated_at is its end time.
        "end_time": row.get("updated_at") if row.get("status") in _TERMINAL_STATUSES else None,
        "sub_jobs": [],
        "slurm_jobs": row.get("slurm_jobs") or [],
        "containers": row.get("containers") or [],
        "logs": row.get("logs") or "",
        "resumed_from_history": True,
    }

def _load_job_for_action(job_id: str, current_user: dict, conn) -> dict:
    """Resolves a job_id to its launch details (genome_path, workflow,
    work_dir, selected_tools, status) from job_store or persisted history.
    Raises 404/403 the same way get_job_status
    does.

    Unlike get_job_status it skips SLURM reconciliation."""
    job = job_store.get(job_id)
    if job is not None:
        if job.get("user_id") != current_user["user_id"]:
            raise HTTPException(status_code=403, detail="Access denied")
        return job

    try:
        user_config = ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn)
    except Exception:
        raise HTTPException(status_code=404, detail="Job not found")
    main_db = user_config.get('main_database')
    row = (
        job_history_client.get_job(
            conn,
            main_db,
            job_id,
            owner_username=current_user["username"],
            owner_cluster_username=current_user["cluster_username"],
        )
        if main_db else None
    )
    if row is None:
        raise HTTPException(status_code=404, detail="Job not found")
    return _job_from_history_row(row)


_TERMINAL_STATUSES = {"completed", "failed", "cancelled"}
_POTENTIALLY_STALE_STATUSES = {"pending", "running", "snakemake"}

# Rows younger than this look interrupted mid-submission; never judged dead.
_SUBMIT_GRACE_SECONDS = 300


def _row_age_seconds(row: dict) -> float | None:
    """Returns seconds since the row was last touched (updated_at, else created_at), or None."""
    stamp = row.get("updated_at") or row.get("created_at")
    if not stamp:
        return None
    try:
        ts = datetime.fromisoformat(stamp)
    except (TypeError, ValueError):
        return None
    if ts.tzinfo is None:
        ts = ts.replace(tzinfo=timezone.utc)
    return (datetime.now(timezone.utc) - ts).total_seconds()


def _within_submit_grace(row: dict) -> bool:
    """Returns True while a row is too young to be judged interrupted."""
    age = _row_age_seconds(row)
    return age is not None and age < _SUBMIT_GRACE_SECONDS
# squeue states that still occupy the cluster (COMPLETING/CONFIGURING count).
_ACTIVE_SLURM_STATES = ("RUNNING", "PENDING", "COMPLETING", "CONFIGURING",
                        "RESIZING", "SUSPENDED", "REQUEUED")

# Stops overlapping polls from reattaching a watcher to the same run.
_REATTACH_LOCK = threading.Lock()


def _has_active_workdir_jobs(conn, work_dir: str, cluster_username: str) -> bool:
    """Returns True if squeue reports any active job under work_dir.

    An SSH or squeue failure counts as alive, so a live run is never marked dead.
    """
    try:
        matches = ssh_slurm.find_active_jobs_in_workdir(work_dir, cluster_username, connection=conn)
    except Exception as exc:
        LOGGER.warning("work_dir SLURM check failed for %s: %s", work_dir, exc)
        return True
    return any(m["state"] in _ACTIVE_SLURM_STATES for m in matches)


def _try_reattach(job_id: str, row: dict, conn, main_db: str | None,
                  current_user: dict) -> dict | None:
    """Resumes watching a detached run after a dane-api restart.

    Replays the job log from ~/.local/share/bsp/jobs/<job_id>.log, which
    re-derives its SLURM jobs, containers and progress. Returns the live job
    dict, or None if the run is not replayable or another thread claimed it.
    """
    # `tail -F` on a log that never grows would block a worker forever.
    try:
        probe = ssh_slurm.probe_run(job_id, connection=conn)
    except Exception as exc:
        LOGGER.warning("Could not probe job %s for reattach: %s", job_id, exc)
        return None
    if not ssh_slurm.is_replayable(probe):
        LOGGER.info("Job %s is not replayable (%s); not reattaching", job_id, probe)
        return None

    with _REATTACH_LOCK:
        if job_store.exists(job_id):
            return job_store.get(job_id)          # another poll won the race

        job_store.create(
            job_id, row.get("genome_path") or "", user_id=current_user["user_id"],
            workflow=row.get("workflow"), output_dir=row.get("work_dir"),
            selected_tools=row.get("selected_tools"),
            relaunched_from=row.get("relaunched_from"),
        )
        # create() would insert a duplicate history row.
        job_store.attach_persistence(job_id, main_db, conn)
        # Carries status/phase over; "pending" would hide the Emergency Stop button.
        job_store.update(job_id, work_dir=row.get("work_dir"),
                         status=row.get("status") or "running",
                         phase=row.get("phase") or "Reattaching to running job")

    LOGGER.info("Reattaching to job %s after an API restart", job_id)
    job_runner.submit_job(job_id, command="", connection=conn, reattach=True)
    return job_store.get(job_id)


@router.get("/job_status/{job_id}")
def get_job_status(
    job_id: str,
    log_offset: int | None = None,
    log_tail: int | None = None,
    current_user: dict = Depends(get_current_user),
):
    """Returns a job's status and log (_job_status).

    log_offset=N returns the log from character N, log_tail=N the last N
    characters, with logs_size the full length; neither returns the whole log."""
    result = _job_status(job_id, current_user)
    if log_offset is None and log_tail is None:
        return result
    full = result.get("logs") or ""
    size = len(full)
    start = max(0, size - log_tail) if log_tail else min(max(0, log_offset or 0), size)
    return {**result, "logs": full[start:], "logs_offset": start, "logs_size": size}


def _job_status(job_id: str, current_user: dict) -> dict:
    """Returns a running job's status, falling back to persistent history.

    Returns 403 if the live job belongs to another user. For a non-terminal
    history row, adds still_active/status_note from squeue.
    """
    job = job_store.get(job_id)
    if job is not None:
        if job.get("user_id") != current_user["user_id"]:
            raise HTTPException(status_code=403, detail="Access denied")
        return {**job, "cluster_host": current_user["cluster_host"]}

    conn = _build_connection(current_user)
    try:
        user_config = ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn)
    except Exception:
        raise HTTPException(status_code=404, detail="Job not found")

    main_db = user_config.get('main_database')
    row = (
        job_history_client.get_job(
            conn,
            main_db,
            job_id,
            owner_username=current_user["username"],
            owner_cluster_username=current_user["cluster_username"],
        )
        if main_db else None
    )
    if row is None:
        raise HTTPException(status_code=404, detail="Job not found")

    # Non-terminal rows lost on restart are reattached by replaying the log.
    if row["status"] not in _TERMINAL_STATUSES:
        live = _try_reattach(job_id, row, conn, main_db, current_user)
        if live is not None:
            return {**live, "cluster_host": current_user["cluster_host"]}

    result = {**_job_from_history_row(row), "cluster_host": current_user["cluster_host"]}

    if row["status"] not in _TERMINAL_STATUSES and row.get("work_dir"):
        try:
            matches = ssh_slurm.find_active_jobs_in_workdir(
                row["work_dir"], current_user["cluster_username"], connection=conn,
            )
        except Exception as exc:
            LOGGER.warning("SLURM reconciliation check failed for job %s: %s", job_id, exc)
            matches = []
        active_matches = [m for m in matches if m["state"] in _ACTIVE_SLURM_STATES]
        still_active = bool(active_matches)
        result["still_active"] = still_active
        if still_active:
            result["status_note"] = (
                "The API was restarted while this job was running. "
                "Phase shown below is from before the restart — live updates will resume shortly."
            )
            if not result.get("slurm_jobs"):
                try:
                    enriched = ssh_slurm.enrich_slurm_jobs_from_logs(
                        row["work_dir"], active_matches, conn,
                    )
                except Exception:
                    enriched = active_matches
                result["slurm_jobs"] = [
                    {
                        "job_id": m["job_id"],
                        "rule": m.get("rule"),
                        "status": m["state"],
                        "time": m.get("time", ""),
                        "genome": m.get("genome"),
                        "source": "fresh run",
                    }
                    for m in enriched
                ]
        elif _within_submit_grace(row):
            # Fresh row with no SLURM jobs yet: still being submitted.
            result["still_active"] = True
            result["status_note"] = (
                "This run is still being submitted to the cluster -- "
                "live updates will appear once its jobs are queued."
            )
        else:
            # Past the grace window with no active jobs: marked interrupted.
            result["status"] = "cancelled"
            result["phase"] = "Interrupted (no active jobs)"
            try:
                job_history_client.record_job_updated(
                    conn,
                    main_db,
                    job_id,
                    status=result["status"],
                    phase=result["phase"],
                )
            except Exception as exc:
                LOGGER.warning("Could not persist stale-status correction for %s: %s", job_id, exc)

    if not result.get("logs") and row.get("work_dir"):
        try:
            result["logs"] = ssh_slurm.read_latest_snakemake_log(row["work_dir"], conn)
        except Exception:
            pass

    return result


def _reconcile_running(conn, main_db: str, rows: list[dict], cluster_username: str) -> None:
    """Corrects rows that claim to be running but are not.

    A run is alive if its .rc sentinel is absent and a driver process exists, its
    driver SLURM job is active, or any worker job under its work_dir is active
    (the driver often exits long before its workers). A present sentinel decides
    the status from its exit code; otherwise the run is interrupted.

    Only rows marked running are touched, and failures are logged and ignored.
    """
    # Everything stays inside the try so a bad row cannot break list_jobs.
    try:
        pending = [r for r in rows
                   if isinstance(r, dict) and (r.get("status") or "").lower() == "running"]
        if not pending:
            return
        ssh = conn.connect()
        for row in pending:
            jid = row.get("job_id") or row.get("id")
            if not jid:
                continue
            # Strict matching: substring matches left finished jobs marked running.
            escaped = re.escape(jid)
            stem = ssh_slurm.run_file_stem(jid)
            # drv: login-node driver process; dj/ds: driver SLURM job id and state.
            probe = (
                f'rc=$(cat $HOME/.local/share/bsp/jobs/{jid}.rc 2>/dev/null); '
                f'drv=$(pgrep -u $USER -f "dane_wf.*{escaped}" 2>/dev/null | wc -l); '
                f'dj=$(cat $HOME/.local/share/bsp/jobs/{stem}.jobid 2>/dev/null); '
                f'ds=$(squeue -h -j "${{dj:-0}}" -o %T 2>/dev/null | head -1); '
                f'echo "${{rc:--}}|$drv|${{ds:--}}"'
            )
            _in, out, _err = ssh.exec_command(probe)
            reply = (out.read().decode() or "").strip().splitlines()
            if not reply:
                continue
            rc, drv, ds = (reply[-1].split("|") + ["-", "0", "-"])[:3]
            try:
                drv_n = int(drv)
            except ValueError:
                continue

            if rc != "-":                      # finished: the sentinel decides
                ok = rc.strip() == "0"
                status, phase = ("completed", "Done") if ok else ("failed", f"Exited {rc.strip()}")
            elif drv_n > 0 or ds.strip() in _ACTIVE_SLURM_STATES:  # genuinely still going
                continue
            elif _within_submit_grace(row):    # too young -- still submitting
                # Mid-submission looks like interrupted; a fresh row stays running.
                continue
            elif row.get("work_dir") and _has_active_workdir_jobs(
                conn, row["work_dir"], cluster_username):
                # Driver exited, but active workers under work_dir mean the run continues.
                continue
            else:                              # no sentinel, no driver, no jobs
                status, phase = "cancelled", "Interrupted (driver stopped)"

            LOGGER.info("reconciled job %s: running -> %s", jid, status)
            row["status"], row["phase"] = status, phase
            try:
                job_history_client.record_job_updated(
                    conn, main_db, jid, status=status, phase=phase)
            except Exception as exc:
                LOGGER.warning("could not persist reconciled status for %s: %s", jid, exc)
    except Exception as exc:
        LOGGER.warning("job reconciliation skipped: %s", exc)


@router.get("/jobs")
def list_jobs(
    workflow: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    current_user: dict = Depends(get_current_user),
):
    """Lists this user's job history, optionally for one workflow, most recent first.
    A user with no history gets an empty list."""
    conn = _build_connection(current_user)
    # The slowest read every page makes: reused for 5 s, then returned while it is asked again.
    return read_cache.get((current_user["user_id"], "jobs", workflow, page, page_size),
                          lambda: _list_jobs(conn, current_user, workflow, page, page_size), fresh=5.0)


def _list_jobs(conn, current_user: dict, workflow: str | None, page: int, page_size: int) -> dict:
    """Reads one page of the job history from the cluster."""
    empty_response = {"jobs": [], "page": page, "page_size": page_size, "total_jobs": 0, "total_pages": 1}
    try:
        user_config = ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn)
    except Exception:
        return empty_response

    main_db = user_config.get('main_database')
    if not main_db:
        return empty_response

    offset = (page - 1) * page_size
    rows, total_jobs = job_history_client.list_jobs_and_count(
        conn,
        main_db,
        workflow=workflow,
        limit=page_size,
        offset=offset,
        owner_username=current_user["username"],
        owner_cluster_username=current_user["cluster_username"],
    )
    # Verifies anything claiming to be running before reporting it.
    _reconcile_running(conn, main_db, rows, current_user["cluster_username"])
    total_pages = max((total_jobs + page_size - 1) // page_size, 1)

    return {
        "jobs": [_job_from_history_row(row) for row in rows],
        "page": page,
        "page_size": page_size,
        "total_jobs": total_jobs,
        "total_pages": total_pages,
    }


@router.post("/cancel_job/{job_id}")
def cancel_job(job_id: str, current_user: dict = Depends(get_current_user)):
    """Emergency stop: cancels the run's SLURM jobs, kills the remote process and marks the job cancelled.

    Falls back to persistent history when the job is not in job_store.
    """
    conn = _build_connection(current_user)
    job = job_store.get(job_id)
    main_db = None

    if job is not None:
        if job.get("user_id") != current_user["user_id"]:
            raise HTTPException(status_code=403, detail="Access denied")
        slurm_ids = [sj["job_id"] for sj in job_store.get_slurm_jobs(job_id)]
    else:
        try:
            user_config = ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn)
        except Exception:
            raise HTTPException(status_code=404, detail="Job not found")

        main_db = user_config.get('main_database')
        row = (
            job_history_client.get_job(
                conn,
                main_db,
                job_id,
                owner_username=current_user["username"],
                owner_cluster_username=current_user["cluster_username"],
            )
            if main_db else None
        )
        if row is None:
            raise HTTPException(status_code=404, detail="Job not found")
        slurm_ids = [sj["job_id"] for sj in (row.get("slurm_jobs") or [])]

    # Driver first; otherwise Snakemake resubmits its children.
    driver = None
    try:
        driver = ssh_slurm.driver_job_id(job_id, connection=conn)
    except Exception as exc:
        LOGGER.warning("Could not look up the driver job for %s: %s", job_id, exc)
    if driver:
        ssh_slurm.cancel_slurm_jobs([driver], connection=conn)
        LOGGER.info("Cancelled driver job %s for job %s", driver, job_id)

    # Cancel all SLURM subjobs
    if slurm_ids:
        ssh_slurm.cancel_slurm_jobs(slurm_ids, connection=conn)
        LOGGER.info("Cancelled %d SLURM jobs for job %s", len(slurm_ids), job_id)

    # Sweeps untracked work_dir jobs; the driver lives in $HOME, so it is untouched.
    work_dir = job.get("work_dir") if job is not None else row.get("work_dir")
    swept_ids: list[str] = []
    if work_dir:
        try:
            stragglers = ssh_slurm.find_active_jobs_in_workdir(
                work_dir, current_user["cluster_username"], connection=conn,
            )
        except Exception as exc:
            LOGGER.warning("Work_dir sweep failed for %s: %s", job_id, exc)
            stragglers = []
        swept_ids = [m["job_id"] for m in stragglers if m["job_id"] not in slurm_ids]
        if swept_ids:
            ssh_slurm.cancel_slurm_jobs(swept_ids, connection=conn)
            LOGGER.info("Cancelled %d untracked work_dir jobs for %s", len(swept_ids), job_id)

    # Only without a driver job: pkill by name would hit other runs.
    if not driver:
        ssh_slurm.kill_remote_process("dane_wf", connection=conn)
        LOGGER.info("Killed remote dane_wf process for job %s", job_id)

    # Cancelling also stops any status checker daemon.
    if job is not None:
        job_store.cancel(job_id)
    else:
        job_history_client.record_job_updated(
            conn, main_db, job_id, status="cancelled", phase="Cancelled by user",
        )

    return {
        "success": True,
        "message": f"Cancelled job {job_id}",
        "slurm_jobs_cancelled": len(slurm_ids) + len(swept_ids)
    }


def _main_db_for(current_user: dict, conn) -> tuple[str, dict]:
    """Returns main_database and the user config, raising HTTPException(400) if either is missing.
    Pre-flight for resume_job/restart_job."""
    try:
        user_config = ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Configuration file not found. Please create a configuration in your Profile settings first.",
        )
    main_db = _resolve_effective_main_db(current_user, conn, user_config, persist=True)
    return main_db, user_config


def _base_output_dir_from(path: str) -> str:
    """Strips a trailing /YYYY-MM-DD-HHMM segment from a job's work_dir, giving the base output dir."""
    return re.sub(r'/\d{4}-\d{2}-\d{2}-\d{4}$', '', path)


@router.post("/resume_job/{job_id}")
def resume_job(job_id: str, current_user: dict = Depends(get_current_user)):
    """Resumes a failed or stale job: copies its work_dir into a new timestamped
    folder and relaunches with margie_sb.resume: true, so completed steps are skipped."""
    _require_current_terms(current_user)
    conn = _build_connection(current_user)
    original = _load_job_for_action(job_id, current_user, conn)

    status = original.get("status")
    work_dir = original.get("work_dir")
    if status not in ("failed", "cancelled"):
        if status not in _POTENTIALLY_STALE_STATUSES:
            raise HTTPException(
                status_code=400,
                detail=f"Cannot resume a job with status '{status}' -- only failed/cancelled jobs, "
                       "or non-terminal jobs no longer actually active on the cluster, can be resumed.",
            )
        # Non-terminal status may be stale after a restart, so the cluster is checked.
        still_active = True
        if work_dir:
            try:
                matches = ssh_slurm.find_active_jobs_in_workdir(
                    work_dir, current_user["cluster_username"], connection=conn,
                )
                still_active = any(m["state"] in ("RUNNING", "PENDING") for m in matches)
            except Exception as exc:
                LOGGER.warning("SLURM active-check failed while resuming job %s: %s", job_id, exc)
                still_active = True  # fail safe: never resume a run not confirmed dead
        if still_active:
            raise HTTPException(
                status_code=400,
                detail=f"Cannot resume a job with status '{status}' while it appears to still be "
                       "active on the cluster.",
            )

    if not work_dir:
        raise HTTPException(status_code=400, detail="Original job has no recorded working directory to resume from")

    genome_path = original.get("genome_path")
    workflow = original.get("workflow")
    if not genome_path or not workflow:
        raise HTTPException(status_code=400, detail="Original job is missing genome_path/workflow, cannot resume")

    main_db, user_config = _main_db_for(current_user, conn)
    slurm_account = str(user_config.get('compute', {}).get('cluster_default', {}).get('account', '')).strip() or None
    slurm_partition = str(user_config.get('compute', {}).get('cluster_default', {}).get('partition', '')).strip() or None
    slurm_walltime = _cluster_default(user_config, 'driver_walltime')
    if not slurm_account:
        raise HTTPException(
            status_code=400,
            detail="compute.cluster_default.account is not configured. Please configure it in your Profile settings.",
        )
    _check_genome_path_exists(genome_path, workflow, conn)

    selected_tools_csv = original.get("selected_tools")
    selected_tools = selected_tools_csv.split(",") if selected_tools_csv else None

    return _launch_job(
        genome_path=genome_path, workflow=workflow,
        base_output_dir=_base_output_dir_from(work_dir), selected_tools=selected_tools,
        current_user=current_user, conn=conn, main_db=main_db,
        slurm_account=slurm_account, slurm_partition=slurm_partition,
        slurm_walltime=slurm_walltime,
        relaunched_from=job_id, copy_from_work_dir=work_dir,
    )


@router.post("/restart_job/{job_id}")
def restart_job(job_id: str, current_user: dict = Depends(get_current_user)):
    """Restarts a job from scratch with the same genome_path/workflow/selected_tools
    in a new timestamped folder. Works from any status."""
    _require_current_terms(current_user)
    conn = _build_connection(current_user)
    original = _load_job_for_action(job_id, current_user, conn)

    base_for_dir = original.get("work_dir") or original.get("output_dir")
    if not base_for_dir:
        raise HTTPException(status_code=400, detail="Original job has no recorded output directory to restart from")

    genome_path = original.get("genome_path")
    workflow = original.get("workflow")
    if not genome_path or not workflow:
        raise HTTPException(status_code=400, detail="Original job is missing genome_path/workflow, cannot restart")

    main_db, user_config = _main_db_for(current_user, conn)
    slurm_account = str(user_config.get('compute', {}).get('cluster_default', {}).get('account', '')).strip() or None
    slurm_partition = str(user_config.get('compute', {}).get('cluster_default', {}).get('partition', '')).strip() or None
    slurm_walltime = _cluster_default(user_config, 'driver_walltime')
    if not slurm_account:
        raise HTTPException(
            status_code=400,
            detail="compute.cluster_default.account is not configured. Please configure it in your Profile settings.",
        )
    _check_genome_path_exists(genome_path, workflow, conn)

    selected_tools_csv = original.get("selected_tools")
    selected_tools = selected_tools_csv.split(",") if selected_tools_csv else None

    return _launch_job(
        genome_path=genome_path, workflow=workflow,
        base_output_dir=_base_output_dir_from(base_for_dir), selected_tools=selected_tools,
        current_user=current_user, conn=conn, main_db=main_db,
        slurm_account=slurm_account, slurm_partition=slurm_partition,
        slurm_walltime=slurm_walltime,
        relaunched_from=job_id,
    )


@router.get("/job_command/{job_id}")
def job_command(job_id: str, current_user: dict = Depends(get_current_user)):
    """The command a job ran, as a script to run again from a shell: its run-command.sh, or, for a
    job from before that file was kept, rebuilt from its record (exact: false; the licence
    settings it ran with were not kept)."""
    conn = _build_connection(current_user)
    job = _load_job_for_action(job_id, current_user, conn)
    work_dir = (job.get("work_dir") or job.get("output_dir") or "").rstrip('/')
    if work_dir:
        path = f"{work_dir}/{COMMAND_FILE}"
        try:
            text = _read_small_remote(path, conn)
            if text.strip():
                return {"script": text, "exact": True, "file": path}
        except OSError:
            pass
        except Exception as exc:
            LOGGER.warning("Could not read %s: %s", path, exc)
    if not job.get("genome_path") or not job.get("workflow"):
        raise HTTPException(status_code=404, detail="No command is recorded for this job")
    return {
        "script": _command_script(
            job_id, job["workflow"], job["genome_path"], work_dir, job.get("selected_tools"),
            note="Rebuilt from the job's record: the licence settings it ran with were not kept.",
        ),
        "exact": False,
        "file": None,
    }


@router.get("/job_files/{job_id}")
def get_job_files(
    job_id: str,
    subdir: str = "",
    current_user: dict = Depends(get_current_user),
):
    """List output files for a job via SFTP."""
    _validate_relative_path(subdir, label="subdirectory")

    conn = _build_connection(current_user)
    work_dir = _resolve_job_work_dir(job_id, current_user, conn)

    target_dir = f"{work_dir}/{subdir}".rstrip("/") if subdir else work_dir

    try:
        entries = ssh_sftp.list_remote_dir(target_dir, connection=conn)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Directory not found on remote")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to list remote directory: {str(e)}")

    return {"work_dir": work_dir, "subdir": subdir, "entries": entries}


@router.get("/download_file/{job_id}")
def download_file(
    job_id: str,
    path: str,
    format: str = Query("raw", pattern="^(raw|excel)$"),
    current_user: dict = Depends(get_current_user),
):
    """Downloads a file from a job's working directory via SFTP.

    format=raw streams the file unmodified; format=excel converts a TSV/CSV
    output to .xlsx in memory.
    """
    _validate_relative_path(path)

    conn = _build_connection(current_user)
    work_dir = _resolve_job_work_dir(job_id, current_user, conn)

    remote_path = f"{work_dir}/{path}"
    filename = path.split("/")[-1]

    if format == "excel":
        xlsx_filename = re.sub(r"\.(tsv|csv)$", "", filename, flags=re.IGNORECASE) + ".xlsx"

        # Pre-generated .xlsx wins; eager read so FileNotFoundError lands here.
        if re.search(r"\.(tsv|csv)$", path, re.IGNORECASE):
            xlsx_remote_path = re.sub(r"\.(tsv|csv)$", ".xlsx", remote_path, flags=re.IGNORECASE)
            try:
                xlsx_bytes = b"".join(ssh_sftp.stream_remote_file(xlsx_remote_path, connection=conn))
                return Response(
                    content=xlsx_bytes,
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{xlsx_filename}"'},
                )
            except (FileNotFoundError, IOError):
                pass  # no pre-generated file — fall through to on-the-fly conversion
            except Exception as e:
                raise HTTPException(status_code=500, detail=f"Failed to read Excel file: {str(e)}")

        try:
            content = b"".join(ssh_sftp.stream_remote_file(remote_path, connection=conn))
        except FileNotFoundError:
            raise HTTPException(status_code=404, detail="File not found on remote")
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Failed to download file: {str(e)}")

        text = content.decode("utf-8", errors="replace")
        delimiter = _detect_delimiter(path, text.splitlines()[0] if text else "")
        try:
            df = pd.read_csv(io.StringIO(text), sep=delimiter, dtype=str, keep_default_na=False)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not parse file as a table: {str(e)}")

        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Sheet1")
            _apply_tier_row_colors(writer.sheets["Sheet1"], df)
        buffer.seek(0)

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{xlsx_filename}"'},
        )

    try:
        return StreamingResponse(
            ssh_sftp.stream_remote_file(remote_path, connection=conn),
            media_type="application/octet-stream",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="File not found on remote")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to download file: {str(e)}")


@router.get("/view_file/{job_id}")
def view_file(
    job_id: str,
    path: str,
    page: int = Query(1, ge=1),
    page_size: int = Query(100, ge=1, le=500),
    current_user: dict = Depends(get_current_user),
):
    """Returns a paginated slice of a delimited text file in a job's work_dir.
    Cost scales with page position, not file size (see ssh_sftp.read_remote_file_page)."""
    _validate_relative_path(path)

    conn = _build_connection(current_user)
    work_dir = _resolve_job_work_dir(job_id, current_user, conn)

    remote_path = f"{work_dir}/{path}"

    cache_key = (job_id, path)
    known_total_lines = None
    try:
        mtime, size = ssh_sftp.stat_remote_file(remote_path, connection=conn)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="File not found on remote")

    cached = _line_count_cache.get(cache_key)
    if cached and cached[0] == mtime and cached[1] == size:
        known_total_lines = cached[2]

    start_row = 2 + (page - 1) * page_size  # row 1 is always the header
    end_row = start_row + page_size - 1

    try:
        result = ssh_sftp.read_remote_file_page(
            remote_path, start_row, end_row, connection=conn,
            known_total_lines=known_total_lines,
        )
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="File not found on remote")
    except IsADirectoryError:
        raise HTTPException(status_code=400, detail="Path is a directory, not a file")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to read file: {str(e)}")

    total_lines = result["total_lines"]
    if known_total_lines is None:
        _line_count_cache[cache_key] = (mtime, size, total_lines)
    total_rows = max(total_lines - 1, 0)
    total_pages = max((total_rows + page_size - 1) // page_size, 1)

    delimiter = _detect_delimiter(path, result["header"])
    columns = result["header"].split(delimiter) if result["header"] else []
    rows = [line.split(delimiter) for line in result["lines"]]

    return {
        "columns": columns,
        "rows": rows,
        "page": page,
        "page_size": page_size,
        "total_rows": total_rows,
        "total_pages": total_pages,
    }


@router.get("/job_status/{job_id}/stream")
def stream_job_status(job_id: str, current_user: dict = Depends(get_current_user)):
    """SSE endpoint that streams real-time job status updates."""
    job = job_store.get(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job not found")
    if job.get("user_id") != current_user["user_id"]:
        raise HTTPException(status_code=403, detail="Access denied")

    return StreamingResponse(
        job_runner.job_status_generator(job_id),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        }
    )


@router.get("/all_genomes")
def all_genomes(path: str, current_user: dict = Depends(get_current_user)):
    """List genome files at a remote path on the user's cluster."""
    conn = _build_connection(current_user)
    genomes = ssh_slurm.get_genomes(path, connection=conn)
    return {"success": True, "Genomes": genomes}


def _resolve_browse_path(path: str, current_user: dict) -> str:
    """Expands a leading ~ to the user's home and normalises the path.

    No allowlist: both /browse endpoints use the user's own SSH credentials, so
    cluster permissions are the boundary (denials surface as 403).
    """
    import posixpath

    if path.startswith("~"):
        path = path.replace("~", current_user["home_dir"], 1)
    return posixpath.normpath(path) if path else "/"


def _permission_status(exc: Exception) -> int:
    """Returns 403 for a permission denial (paramiko raises OSError with EACCES), else 500."""
    import errno as _errno
    if isinstance(exc, PermissionError):
        return 403
    if isinstance(exc, OSError) and exc.errno == _errno.EACCES:
        return 403
    if "permission denied" in str(exc).lower():
        return 403
    return 500


@router.get("/browse")
def browse(path: str, current_user: dict = Depends(get_current_user)):
    """Lists a remote directory for the file explorer: directories first, then
    files, alphabetically, plus the parent path for breadcrumbs.

    Runs with the user's own SSH credentials, so no traversal guard is needed.
    """
    import posixpath

    conn = _build_connection(current_user)
    path = _resolve_browse_path(path, current_user)

    try:
        entries = ssh_sftp.list_remote_dir_checked(path, connection=conn)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail=f"Path not found on cluster: '{path}'")
    except NotADirectoryError:
        raise HTTPException(status_code=400, detail=f"Not a directory: '{path}'")
    except Exception as exc:
        status = _permission_status(exc)
        detail = f"Permission denied: '{path}'" if status == 403 else f"Failed to list directory: {exc}"
        raise HTTPException(status_code=status, detail=detail)

    entries.sort(key=lambda e: (e["type"] != "directory", e["name"].lower()))
    parent = posixpath.dirname(path.rstrip("/")) or "/"

    return {
        "success": True,
        "path": path,
        "parent": parent,
        "entries": entries,
    }


_FINAL_SUMMARIES: dict = {}


@router.get("/final_summary")
def final_summary(path: str, current_user: dict = Depends(get_current_user)):
    """Returns a FINAL confidence table's report numbers (services/final_summary.py).

    Read from disk when this API runs on the cluster as the user, else over SFTP;
    cached until the file's size or mtime changes."""
    import os
    from bioinformatics_tools.api.services import final_summary as fs
    from bioinformatics_tools.utilities.ssh_connection import runs_here

    conn = _build_connection(current_user)
    path = _resolve_browse_path(path, current_user)
    if not path.endswith(".tsv"):
        raise HTTPException(status_code=400, detail="Not a FINAL table (.tsv)")
    try:
        if runs_here(conn) and os.path.exists(path):
            st = os.stat(path)
            key = (current_user["user_id"], path, st.st_size, st.st_mtime)
            if key not in _FINAL_SUMMARIES:
                with open(path, encoding="utf-8", errors="replace") as fh:
                    _FINAL_SUMMARIES[key] = fs.tally(fh)
        else:
            sftp = conn.connect().open_sftp()
            try:
                st = sftp.stat(path)
                key = (current_user["user_id"], path, st.st_size, st.st_mtime)
                if key not in _FINAL_SUMMARIES:
                    with sftp.open(path, "r") as fh:
                        fh.prefetch()
                        _FINAL_SUMMARIES[key] = fs.tally(l.decode("utf-8", "replace") if isinstance(l, bytes) else l for l in fh)
            finally:
                sftp.close()
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail=f"Path not found on cluster: '{path}'")
    except Exception as exc:
        raise HTTPException(status_code=_permission_status(exc), detail=f"Could not read {path}: {exc}")
    while len(_FINAL_SUMMARIES) > 64:
        _FINAL_SUMMARIES.pop(next(iter(_FINAL_SUMMARIES)))
    return _FINAL_SUMMARIES[key]


# Byte cap for the in-browser file viewer; the UI flags truncation.
_VIEW_MAX_BYTES = 1_000_000


@router.get("/browse_view")
def browse_view(path: str, current_user: dict = Depends(get_current_user)):
    """Returns the head of a text file on the cluster for the explorer's View button.

    Reads at most _VIEW_MAX_BYTES and refuses binary files (a NUL byte).
    """
    conn = _build_connection(current_user)
    path = _resolve_browse_path(path, current_user)

    try:
        kind = ssh_sftp.check_remote_path_kind(path, conn)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail=f"File not found on cluster: '{path}'")
    except Exception as exc:
        raise HTTPException(status_code=_permission_status(exc),
                            detail=f"Could not access path: {exc}")

    if kind == "directory":
        raise HTTPException(status_code=400, detail=f"Path is a directory, not a file: '{path}'")

    try:
        chunks, total = [], 0
        for chunk in ssh_sftp.stream_remote_file(path, connection=conn):
            chunks.append(chunk)
            total += len(chunk)
            if total >= _VIEW_MAX_BYTES:
                break
        raw = b"".join(chunks)[:_VIEW_MAX_BYTES]
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="File not found on remote")
    except Exception as exc:
        status = _permission_status(exc)
        detail = f"Permission denied: '{path}'" if status == 403 else f"Failed to read file: {exc}"
        raise HTTPException(status_code=status, detail=detail)

    if b"\x00" in raw:
        return {
            "path": path,
            "binary": True,
            "truncated": False,
            "content": "",
        }

    return {
        "path": path,
        "binary": False,
        "truncated": total >= _VIEW_MAX_BYTES,
        "content": raw.decode("utf-8", errors="replace"),
    }


def _parse_path_size(out: str) -> dict:
    """Reads the size probe's output: a kind letter, the bytes, then the file count."""
    parts = out.split()
    if not parts or parts[0] not in ("d", "f"):
        raise FileNotFoundError
    nums = [int(p) for p in parts[1:3] if p.isdigit()]
    return {
        "kind": "directory" if parts[0] == "d" else "file",
        "bytes": nums[0] if nums else None,
        "files": nums[1] if len(nums) > 1 else None,
    }


@router.get("/path_size")
def path_size(path: str, current_user: dict = Depends(get_current_user)):
    """Returns a cluster path's kind, size in bytes (a folder's total) and file count, so the
    Files page can ask before a large download."""
    conn = _build_connection(current_user)
    path = _resolve_browse_path(path, current_user)
    q = shlex.quote(path)
    _, out = _run_remote_check(
        conn,
        f"if [ -d {q} ]; then echo d; du -sb {q} 2>/dev/null | cut -f1; find {q} -type f 2>/dev/null | wc -l; "
        f"elif [ -e {q} ]; then echo f; stat -Lc %s {q}; echo 1; else echo missing; fi",
    )
    try:
        return {"path": path, **_parse_path_size(out)}
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail=f"Path not found on cluster: '{path}'")


def _archive_command(parent: str, name: str, has_zip: bool) -> tuple[str, str, str]:
    """The command that writes a folder as an archive to stdout, its file name and media type:
    zip where the cluster has it, else a gzipped tar."""
    if has_zip:
        return f"cd {shlex.quote(parent)} && zip -r -q - {shlex.quote(name)} 2>/dev/null", f"{name}.zip", "application/zip"
    return f"tar -C {shlex.quote(parent)} -czf - {shlex.quote(name)} 2>/dev/null", f"{name}.tar.gz", "application/gzip"


def _attachment(name: str) -> dict:
    safe = name.replace('"', "").replace("\\", "") or "download"
    return {"Content-Disposition": f'attachment; filename="{safe}"'}


@router.get("/download_path")
def download_path(path: str, current_user: dict = Depends(get_current_user)):
    """Streams any file the account can read, whole; a folder is zipped on the cluster as it streams.

    Unlike download_file, the path need not be inside a job's folder."""
    conn = _build_connection(current_user)
    path = _resolve_browse_path(path, current_user)
    try:
        kind = ssh_sftp.check_remote_path_kind(path, conn)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail=f"Path not found on cluster: '{path}'")
    except Exception as exc:
        raise HTTPException(status_code=_permission_status(exc), detail=f"Could not access path: {exc}")

    name = posixpath.basename(path.rstrip("/")) or "cluster"
    if kind != "directory":
        try:
            chunks = ssh_sftp.stream_remote_file(path, connection=conn)
        except FileNotFoundError:
            raise HTTPException(status_code=404, detail="File not found on remote")
        except Exception as exc:
            raise HTTPException(status_code=_permission_status(exc), detail=f"Failed to read file: {exc}")
        return StreamingResponse(chunks, media_type="application/octet-stream", headers=_attachment(name))

    has_zip = _run_remote_check(conn, "command -v zip >/dev/null 2>&1")[0] == 0
    command, filename, media = _archive_command(posixpath.dirname(path.rstrip("/")) or "/", name, has_zip)
    _, stdout, _ = conn.connect().exec_command(command)

    def _chunks():
        try:
            while True:
                block = stdout.read(65536)
                if not block:
                    break
                yield block
        finally:
            stdout.channel.close()

    return StreamingResponse(_chunks(), media_type=media, headers=_attachment(filename))


# ---- Copy and move on the cluster (Files): a job of its own, on SLURM or the login node ----

# Where each transfer keeps its plan, script, status and results, under the user's home.
TRANSFER_DIR = ".margie/transfers"
# Cores for a transfer on SLURM (items copy in parallel); on the login node, fewer.
TRANSFER_CORES = 8
LOGIN_CORES = 2

# Runs the plan (op \t source \t target per line), a few items at a time. It never overwrites:
# a target that exists is reported, not replaced; a delete has no target. Each item's result
# goes in r.<n>, then all in results.
TRANSFER_SCRIPT = r"""#!/bin/bash
set -u
dir="$1"; n="${2:-2}"
say() { printf 'state=%s\nhost=%s\n' "$1" "$(hostname)" > "$dir/status.tmp" && mv "$dir/status.tmp" "$dir/status"; }
say running
i=0
while IFS=$'\t' read -r op src dst; do
  i=$((i+1))
  while [ "$(jobs -rp | wc -l)" -ge "$n" ]; do wait -n; done
  (
    if [ ! -e "$src" ] && [ ! -L "$src" ]; then r=MISSING
    elif [ "$op" = delete ]; then out=$(rm -rf -- "$src" 2>&1) && r=OK || r="ERR ${out//$'\n'/ }"
    elif [ -e "$dst" ] || [ -L "$dst" ]; then r=EXISTS
    elif [ "$op" = copy ]; then out=$(cp -Rp -- "$src" "$dst" 2>&1) && r=OK || r="ERR ${out//$'\n'/ }"
    else out=$(mv -- "$src" "$dst" 2>&1) && r=OK || r="ERR ${out//$'\n'/ }"
    fi
    printf '%s\t%s\n' "$i" "$r" > "$dir/r.$i"
  ) &
done < "$dir/plan"
wait
cat "$dir"/r.* 2>/dev/null | sort -n > "$dir/results"
say done
"""

_RESULT = {
    "MISSING": "It is no longer there.",
    "EXISTS": "Something of that name is already in this folder; nothing was replaced.",
}


def _transfer_plan(op: str, sources: list[str], into: str) -> tuple[list[dict], list[dict]]:
    """Splits the sources into items to copy or move into `into`, and those refused at once:
    a folder into itself, an item already in the folder, two items of one name, or a name
    the plan cannot hold (a tab or a line break)."""
    items, refused, names = [], [], set()
    for source in sources:
        target = posixpath.join(into, posixpath.basename(source))
        name = posixpath.basename(source)
        why = ""
        if source == "/" or into == source or into.startswith(source + "/"):
            why = "A folder cannot go inside itself."
        elif posixpath.dirname(source) == into:
            why = "It is already in this folder."
        elif name in names:
            why = "Another item of the same name is in this paste; nothing was replaced."
        elif any(c in source for c in "\t\n\r"):
            why = "Its name has a tab or a line break, which cannot be copied from here."
        if why:
            refused.append({"source": source, "target": target, "ok": False, "message": why})
        else:
            names.add(name)
            items.append({"source": source, "target": target})
    return items, refused


def _transfer_results(items: list[dict], raw: str) -> list[dict]:
    """Each item's outcome from the results lines (n \t OK | MISSING | EXISTS | ERR message)."""
    got = {}
    for line in raw.splitlines():
        n, _, r = line.partition("\t")
        if n.isdigit():
            got[int(n)] = r
    out = []
    for k, it in enumerate(items, start=1):
        r = got.get(k)
        if r is None:
            out.append({**it, "ok": None, "message": ""})
        elif r == "OK":
            out.append({**it, "ok": True, "message": ""})
        else:
            out.append({**it, "ok": False, "message": _RESULT.get(r, r[4:] if r.startswith("ERR ") else r)})
    return out


@router.post("/transfer_paths")
def transfer_paths(payload: dict, current_user: dict = Depends(get_current_user)):
    """Copies or moves files and folders into a folder on the cluster, as the user, as a job:
    on SLURM (the user's account and main partition, 8 cores) or, when SLURM will not take it,
    detached on the login node. Nothing is overwritten. Follow it with GET /transfer_paths/{id}.

    Body: {"op": "copy" | "move", "sources": [paths], "into": folder}."""
    op = payload.get("op")
    sources = payload.get("sources")
    into = payload.get("into")
    if op not in ("copy", "move"):
        raise HTTPException(status_code=400, detail="op must be copy or move")
    if not isinstance(sources, list) or not sources or len(sources) > 500 or not all(isinstance(x, str) and x for x in sources):
        raise HTTPException(status_code=400, detail="sources must be a list of 1 to 500 paths")
    if not isinstance(into, str) or not into:
        raise HTTPException(status_code=400, detail="into must be a folder")

    conn = _build_connection(current_user)
    into = _resolve_browse_path(into, current_user).rstrip("/") or "/"
    code, _ = _run_remote_check(conn, f"test -d {shlex.quote(into)}")
    if code != 0:
        raise HTTPException(status_code=404, detail=f"Not a folder on the cluster: '{into}'")
    resolved = [_resolve_browse_path(x, current_user).rstrip("/") or "/" for x in sources]
    items, refused = _transfer_plan(op, resolved, into)
    if not items:
        return {"id": None, "op": op, "into": into, "state": "done", "where": None, "job": None, "results": refused}

    tid = uuid.uuid4().hex[:12]
    d = f"{current_user['home_dir']}/{TRANSFER_DIR}/{tid}"
    q = shlex.quote
    code, out = _run_remote_check(conn, f"mkdir -p {q(d)} && chmod 700 {q(d)}")
    if code != 0:
        raise HTTPException(status_code=500, detail=f"Could not make {d}: {out}")
    ssh_sftp.write_remote_text_file(f"{d}/run.sh", TRANSFER_SCRIPT, connection=conn)
    ssh_sftp.write_remote_text_file(f"{d}/plan", "".join(f"{op}\t{it['source']}\t{it['target']}\n" for it in items), connection=conn)
    ssh_sftp.write_remote_text_file(f"{d}/meta.json", json.dumps({"op": op, "into": into, "items": items, "refused": refused}) + "\n", connection=conn)

    try:
        cfg = ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn) or {}
    except Exception:
        cfg = {}
    compute = (cfg.get("compute") or {}).get("cluster_default") or {}
    account = str(compute.get("account") or "").strip()
    partition = str(compute.get("partition") or "").strip()
    where, job = "login", None
    if account:
        code, out = _run_remote_check(
            conn,
            f"cd {q(d)} && sbatch --parsable --job-name=margie-{op} --account={q(account)} "
            + (f"--partition={q(partition)} " if partition else "")
            + f"--time=04:00:00 --ntasks=1 --cpus-per-task={TRANSFER_CORES} --output={q(d)}/slurm.out "
            + f"run.sh {q(d)} {TRANSFER_CORES}",
        )
        first = out.strip().splitlines()[-1].split(";")[0].strip() if out.strip() else ""
        if code == 0 and first.isdigit():
            where, job = "slurm", first
            ssh_sftp.write_remote_text_file(f"{d}/status", "state=queued\n", connection=conn)
            ssh_sftp.write_remote_text_file(f"{d}/job", job + "\n", connection=conn)
    if where == "login":
        # SLURM would not take it (or no account is set): the login node does it.
        code, out = _run_remote_check(conn, f"cd {q(d)} && setsid nohup bash run.sh {q(d)} {LOGIN_CORES} > /dev/null 2>&1 < /dev/null & echo started")
        if "started" not in out:
            raise HTTPException(status_code=500, detail=f"Could not start the {op}: {out}")
    return {"id": tid, "op": op, "into": into, "state": "queued" if job else "running", "where": where, "job": job,
            "results": refused + [{**it, "ok": None, "message": ""} for it in items]}


@router.get("/transfer_paths/{tid}")
def transfer_status(tid: str, current_user: dict = Depends(get_current_user)):
    """A transfer's state (queued, running, done or failed) and each item's outcome so far."""
    if not re.fullmatch(r"[0-9a-f]{12}", tid):
        raise HTTPException(status_code=400, detail="Not a transfer id")
    conn = _build_connection(current_user)
    d = f"{current_user['home_dir']}/{TRANSFER_DIR}/{tid}"
    q = shlex.quote
    _, out = _run_remote_check(
        conn,
        f"cd {q(d)} 2>/dev/null || exit 3; echo '@@meta'; cat meta.json; echo '@@status'; cat status 2>/dev/null; "
        f"echo '@@job'; cat job 2>/dev/null; echo '@@results'; cat results 2>/dev/null || cat r.* 2>/dev/null; echo '@@end'",
    )
    if "@@meta" not in out:
        raise HTTPException(status_code=404, detail="No such transfer")
    part = lambda name: out.split(f"@@{name}", 1)[1].split("@@", 1)[0].strip()  # noqa: E731
    meta = json.loads(part("meta"))
    state = dict(line.split("=", 1) for line in part("status").splitlines() if "=" in line).get("state", "running")
    job = part("job") or None
    if state == "queued" and job:
        # A job gone from the queue without starting failed (no account, wrong partition, ...).
        _, squeue = _run_remote_check(conn, f"squeue -h -j {q(job)} -o %T 2>/dev/null")
        if not squeue.strip():
            _, status_now = _run_remote_check(conn, f"cat {q(d)}/status 2>/dev/null")
            if "state=queued" in status_now:
                state = "failed"
    results = meta["refused"] + _transfer_results(meta["items"], part("results"))
    return {"id": tid, "op": meta["op"], "into": meta["into"], "state": state, "where": "slurm" if job else "login", "job": job, "results": results}


# What MARGIE keeps in the home folder (its code, keys, settings, job records): never deleted
# from here, nor anything in them or above them.
_KEEP_IN_HOME = ("bioinformatics-tools", ".ssh", ".config/bioinformatics-tools", ".local/share/bsp", ".margie")


def _delete_plan(sources: list[str], seen: dict[str, str], home: str, code: str, group: str) -> tuple[list[dict], list[dict]]:
    """Splits the paths into those to delete and those refused at once. `seen` maps each path
    to where it really is (symlinked folders above it followed); `code` is where MARGIE's code
    really is and `group` the lab's shared MARGIE folder ('' when none). Refused: the top of
    the file system, home or anything above it, MARGIE's own folders (_KEEP_IN_HOME, its code),
    the shared folder itself, and its databases/ and sif/, which every run uses."""
    above = lambda p, q: q == p or q.startswith(p.rstrip("/") + "/")  # noqa: E731  (p is q or above it)
    whole = [f"{home}/{k}" for k in _KEEP_IN_HOME] + [x for x in (code,) if x]
    shared = [f"{group}/databases", f"{group}/sif"] if group else []
    items, refused, done = [], [], set()
    for source in sources:
        why = ""
        for p in {source, seen.get(source) or source}:
            if p.count("/") < 2 or any(c in p for c in "\t\n\r"):
                why = why or "Too high up in the file system to delete from here."
            elif above(p, home):
                why = why or "Your home folder, or a folder above it, is not deleted from here."
            elif any(above(p, k) or above(k, p) for k in whole):
                why = why or "MARGIE's own code, keys, settings or job records: they are not deleted from here."
            elif any(above(p, k) or above(k, p) for k in shared) or (group and above(p, group)):
                why = why or "The lab's shared databases and containers, which every run uses, are not deleted from here."
        if why:
            refused.append({"source": source, "target": "", "ok": False, "message": why})
        elif source not in done:
            done.add(source)
            items.append({"source": source, "target": ""})
    return items, refused


@router.post("/delete_paths")
def delete_paths(payload: dict, current_user: dict = Depends(get_current_user)):
    """Deletes files and folders (and everything in them) on the cluster, as the user, as a job
    detached on the login node, so a large folder may take its time; where the account may not
    delete, the cluster refuses. Some paths are never deleted (_delete_plan). Follow it with
    GET /transfer_paths/{id}.

    Body: {"paths": [paths]}."""
    paths = payload.get("paths")
    if not isinstance(paths, list) or not paths or len(paths) > 500 or not all(isinstance(x, str) and x.strip() for x in paths):
        raise HTTPException(status_code=400, detail="paths must be a list of 1 to 500 paths")
    conn = _build_connection(current_user)
    home = posixpath.normpath(current_user["home_dir"])
    resolved = [_resolve_browse_path(x.strip(), current_user).rstrip("/") or "/" for x in paths]
    q = shlex.quote
    try:
        cfg = ssh_sftp.read_remote_yaml(_config_path(current_user["home_dir"]), connection=conn) or {}
    except Exception:
        cfg = {}
    group = user_stores.depot(cfg)
    # Where MARGIE's code and each path really are, symlinked folders followed: one @line each.
    _, out = _run_remote_check(
        conn,
        f'printf "@%s\\n" "$(readlink -f -- {q(home + "/bioinformatics-tools")} 2>/dev/null)"; '
        + "; ".join(f'd=$(readlink -f -- "$(dirname -- {q(p)})" 2>/dev/null) && printf "@%s/%s\\n" "$d" "$(basename -- {q(p)})" || echo @' for p in resolved),
    )
    lines = [line[1:].strip() for line in out.splitlines() if line.startswith("@")]
    code = lines[0] if lines else ""
    seen = {p: posixpath.normpath(lines[k + 1]) for k, p in enumerate(resolved) if k + 1 < len(lines) and lines[k + 1]}
    items, refused = _delete_plan(resolved, seen, home, code, group)
    if not items:
        return {"id": None, "op": "delete", "into": "", "state": "done", "where": None, "job": None, "results": refused}

    tid = uuid.uuid4().hex[:12]
    d = f"{current_user['home_dir']}/{TRANSFER_DIR}/{tid}"
    code_, out = _run_remote_check(conn, f"mkdir -p {q(d)} && chmod 700 {q(d)}")
    if code_ != 0:
        raise HTTPException(status_code=500, detail=f"Could not make {d}: {out}")
    ssh_sftp.write_remote_text_file(f"{d}/run.sh", TRANSFER_SCRIPT, connection=conn)
    ssh_sftp.write_remote_text_file(f"{d}/plan", "".join(f"delete\t{it['source']}\t\n" for it in items), connection=conn)
    ssh_sftp.write_remote_text_file(f"{d}/meta.json", json.dumps({"op": "delete", "into": "", "items": items, "refused": refused}) + "\n", connection=conn)
    _, out = _run_remote_check(conn, f"cd {q(d)} && setsid nohup bash run.sh {q(d)} {LOGIN_CORES} > /dev/null 2>&1 < /dev/null & echo started")
    if "started" not in out:
        raise HTTPException(status_code=500, detail=f"Could not start deleting: {out}")
    return {"id": tid, "op": "delete", "into": "", "state": "running", "where": "login", "job": None,
            "results": refused + [{**it, "ok": None, "message": ""} for it in items]}


# Largest file the advanced editor saves (config.yaml, the Snakemake workflow).
_EDIT_MAX_BYTES = 2_000_000


@router.post("/edit_file")
def edit_file(payload: dict, current_user: dict = Depends(get_current_user)):
    """Saves a text file edited by hand (Settings → Advanced: config.yaml, the Snakemake workflow).

    Body: {"path", "content", "check": "yaml" | null}. With check "yaml" the text must parse as a
    YAML mapping, or nothing is written. The version it replaces is kept beside it as
    <name>.bak-<date-time>, so an edit can always be undone."""
    path = payload.get("path")
    content = payload.get("content")
    check = payload.get("check")
    if not isinstance(path, str) or not path.strip():
        raise HTTPException(status_code=400, detail="Which file?")
    if not isinstance(content, str) or len(content.encode()) > _EDIT_MAX_BYTES:
        raise HTTPException(status_code=400, detail="The text is missing or larger than 2 MB.")
    if check == "yaml":
        try:
            parsed = yaml.safe_load(content)
        except yaml.YAMLError as exc:
            mark = getattr(exc, "problem_mark", None)
            where = f" (line {mark.line + 1}, column {mark.column + 1})" if mark else ""
            raise HTTPException(status_code=400, detail=f"This is not valid YAML{where}: {getattr(exc, 'problem', None) or exc}. Nothing was saved.")
        if not isinstance(parsed, dict):
            raise HTTPException(status_code=400, detail="The configuration must be a YAML mapping (key: value lines). Nothing was saved.")

    conn = _build_connection(current_user)
    path = _resolve_browse_path(path.strip(), current_user)
    q = shlex.quote
    backup = f"{path}.bak-{time.strftime('%Y%m%d-%H%M%S')}"
    code, out = _run_remote_check(conn, f"if [ -e {q(path)} ]; then cp -p -- {q(path)} {q(backup)} && echo KEPT; else echo NEW; fi")
    if code != 0:
        raise HTTPException(status_code=_permission_status(Exception(out)), detail=f"Could not keep a copy of {path}: {out}")
    try:
        ssh_sftp.write_remote_text_file(path, content, connection=conn)
    except Exception as exc:
        raise HTTPException(status_code=_permission_status(exc), detail=f"Could not save {path}: {exc}")
    return {"path": path, "backup": backup if "KEPT" in out else None}


@router.post("/make_path")
def make_path(payload: dict, current_user: dict = Depends(get_current_user)):
    """Makes a folder, with any missing folders above it (mkdir -p), or a new empty text file,
    as the user: where their account may not write (someone else's folders) the cluster
    refuses, and so does this. Nothing that exists is replaced.

    Body: {"path", "kind": "folder" | "file"}."""
    path = payload.get("path")
    kind = payload.get("kind")
    if kind not in ("folder", "file"):
        raise HTTPException(status_code=400, detail="kind must be folder or file")
    if not isinstance(path, str) or not path.strip() or any(c in path for c in "\n\r\t"):
        raise HTTPException(status_code=400, detail="Which path?")
    conn = _build_connection(current_user)
    path = _resolve_browse_path(path.strip(), current_user)
    q = shlex.quote(path)
    if kind == "folder":
        command = f"if [ -d {q} ]; then echo THERE; elif [ -e {q} ]; then echo NOT_A_FOLDER; else mkdir -p -- {q} 2>&1 && echo MADE; fi"
    else:
        parent = shlex.quote(posixpath.dirname(path) or "/")
        command = (f"if [ -e {q} ] || [ -L {q} ]; then echo EXISTS; elif [ ! -d {parent} ]; then echo NO_FOLDER; "
                   f"else (set -o noclobber; : > {q}) 2>&1 && echo MADE; fi")
    _, out = _run_remote_check(conn, command)
    last = out.strip().splitlines()[-1] if out.strip() else ""
    if last == "MADE":
        return {"path": path, "kind": kind, "made": True}
    if last == "THERE":
        return {"path": path, "kind": kind, "made": False}
    if last == "EXISTS":
        raise HTTPException(status_code=409, detail=f"{posixpath.basename(path)} is already there; nothing was replaced.")
    if last == "NOT_A_FOLDER":
        raise HTTPException(status_code=409, detail=f"{path} is a file, not a folder.")
    if last == "NO_FOLDER":
        raise HTTPException(status_code=404, detail=f"The folder {posixpath.dirname(path)} does not exist.")
    if "ermission denied" in out:
        raise HTTPException(status_code=403, detail=f"Your account may not create {path}: permission denied there.")
    raise HTTPException(status_code=500, detail=f"Could not create {path}: {out.strip() or 'no answer from the cluster'}")


@router.post("/browse_save")
def browse_save(payload: dict, current_user: dict = Depends(get_current_user)):
    """Writes edited text back to a file on the cluster from the explorer's editor.

    Refuses directories and truncated content, which would discard the unread part.
    """
    path = payload.get("path", "").strip()
    content = payload.get("content", "")
    if not path:
        raise HTTPException(status_code=400, detail="Path is required")
    if payload.get("truncated"):
        raise HTTPException(status_code=400, detail="File is too large to edit in-browser")

    conn = _build_connection(current_user)
    resolved_path = _resolve_browse_path(path, current_user)

    try:
        kind = ssh_sftp.check_remote_path_kind(resolved_path, conn)
        if kind == "directory":
            raise HTTPException(status_code=400, detail=f"Path is a directory, not a file: '{resolved_path}'")
    except FileNotFoundError:
        pass  # new file — fine to create on save

    try:
        ssh_sftp.write_remote_text_file(resolved_path, content, connection=conn)
    except Exception as exc:
        status = _permission_status(exc)
        detail = f"Permission denied: '{resolved_path}'" if status == 403 else f"Failed to save file: {exc}"
        raise HTTPException(status_code=status, detail=detail)

    return {"success": True, "path": resolved_path}
